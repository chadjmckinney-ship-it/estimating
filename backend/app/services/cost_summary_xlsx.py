"""
The estimate Summary as an .xlsx (sql/088): the workbook's Summary tab.

One sheet, laid out the way the tab is: the job's header, the JOB COST
SUMMARY (a row per section, a $/unit line under it, the totals, each
column's share of the price), the lower block (TOTAL MATERIAL / LABOR / SUB
LABOR / OTHER, sales tax, Margin Contingency, ESTIMATED PROFIT, the labor
insurance the tab notes), then the COST CODES with a Summary column and a
column per section. The section figures are values; the totals, the $/unit
lines, the shares and the lower block are formulas, so a number nudged in
Excel still adds up — the proposal form's rule.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from app.services.cost_summary import COLUMNS, LABOR_COLUMNS, MATERIAL_COLUMNS, OTHER_COLUMNS, SUB_LABOR_COLUMNS

ARIAL = "Arial"
DARK, GREY, NAVY, LIGHT, WHITE = "1C1C1C", "5C6570", "1C3F8A", "E8EEF8", "FFFFFF"
MONEY0 = '"$"#,##0'
MONEY2 = '"$"#,##0.00'
PCT = "0.0%"
QTY = "#,##0.###"
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
SUMMARY_TOTAL_ROW_LABEL = "TOTALS"


def _font(size: float = 9, bold: bool = False, color: str = DARK, italic: bool = False) -> Font:
    return Font(name=ARIAL, size=size, bold=bold, color=color, italic=italic)


def _fill(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


def _num(x: Any) -> float | None:
    if x is None or x == "":
        return None
    return float(Decimal(str(x)))


class _Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.r = 0

    def row(self, height: float | None = None) -> int:
        self.r += 1
        if height is not None:
            self.ws.row_dimensions[self.r].height = height
        return self.r

    def put(self, col: int, value: Any, *, font: Font | None = None, fill: PatternFill | None = None,
            align: Alignment | None = None, border: Border | None = None, fmt: str | None = None):
        c = self.ws.cell(self.r, col)
        c.value = value
        c.font = font or _font()
        if fill is not None:
            c.fill = fill
        if align is not None:
            c.alignment = align
        if border is not None:
            c.border = border
        if fmt is not None:
            c.number_format = fmt
        return c


def _visible_columns(s: dict[str, Any]) -> list[tuple[str, str]]:
    """The tab's fifteen, the tax, and UNASSIGNED only when something is."""
    unassigned = Decimal(str(s["totals"]["columns"].get("unassigned", 0) or 0))
    return [(k, label) for k, label in COLUMNS if k != "unassigned" or unassigned != 0]


def build_workbook(s: dict[str, Any]) -> bytes:
    """The summary read (services.cost_summary.estimate_summary) as the tab, in bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    sh = _Sheet(ws)
    cols = _visible_columns(s)
    sections = s["sections"]
    first_money_col = 6                      # F: the tab's G, one over for the $/unit column
    letter = {k: get_column_letter(first_money_col + i) for i, (k, _) in enumerate(cols)}
    last_col = first_money_col + len(cols) - 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 11
    for i in range(len(cols)):
        ws.column_dimensions[get_column_letter(first_money_col + i)].width = 13

    # ------------------------------------------------------------ header --
    sh.row(22)
    sh.put(1, "JOB COST SUMMARY", font=_font(14, True, NAVY))
    for label, value in (
        ("JOB NAME:", s.get("project_name")),
        ("JOB ADDRESS:", s.get("location")),
        ("BID #:", s.get("job_number")),
        ("GC:", s.get("gc")),
        ("ESTIMATE:", f"{s.get('estimate_name') or ''}  ({s.get('estimate_status') or ''})"),
        ("ESTIMATOR:", s.get("estimator_name")),
        ("DATE:", s["generated_at"].strftime("%Y-%m-%d") if hasattr(s["generated_at"], "strftime") else str(s["generated_at"])),
    ):
        sh.row(14)
        sh.put(1, label, font=_font(9, True, GREY))
        sh.put(2, value or "", font=_font(9))
    sh.row(8)

    # ------------------------------------------------- job cost summary --
    sh.row(30)
    heads = [("DESCRIPTION", LEFT), ("QUANTITY", CENTER), ("UNT", CENTER), ("SALE", CENTER), ("$ / UNIT", CENTER)]
    for i, (text, align) in enumerate(heads, start=1):
        sh.put(i, text, font=_font(8, True, WHITE), fill=_fill(NAVY), align=align, border=BOX)
    for k, label in cols:
        sh.put(first_money_col + cols.index((k, label)), label, font=_font(8, True, WHITE), fill=_fill(NAVY),
               align=CENTER, border=BOX)

    money_rows: list[int] = []
    for sec in sections:
        r = sh.row(16)
        money_rows.append(r)
        sh.put(1, sec["name"], font=_font(9, True), border=BOX)
        sh.put(2, _num(sec.get("quantity")), font=_font(9), border=BOX, fmt=QTY, align=RIGHT)
        sh.put(3, sec.get("unit") or "", font=_font(9), border=BOX, align=CENTER)
        sh.put(4, _num(sec["sale"]), font=_font(9, True), border=BOX, fmt=MONEY0, align=RIGHT)
        sh.put(5, f"=IF(B{r}=0,0,D{r}/B{r})", font=_font(9), border=BOX, fmt=MONEY2, align=RIGHT)
        for k, _ in cols:
            sh.put(first_money_col + cols.index((k, _)), _num(sec["columns"].get(k, 0)), font=_font(9),
                   border=BOX, fmt=MONEY0, align=RIGHT)
        # The tab's "$ /SF" line under each section.
        r2 = sh.row(13)
        sh.put(1, f"$ / {sec.get('unit') or 'unit'}", font=_font(8, color=GREY, italic=True), align=RIGHT)
        for k, _ in cols:
            sh.put(first_money_col + cols.index((k, _)), f"=IF($B{r}=0,0,{letter[k]}{r}/$B{r})",
                   font=_font(8, color=GREY, italic=True), fmt=MONEY2, align=RIGHT)
        _ = r2

    tot = sh.row(18)
    sh.put(1, SUMMARY_TOTAL_ROW_LABEL, font=_font(9, True, WHITE), fill=_fill(NAVY), border=BOX)
    for col in (2, 3, 5):
        sh.put(col, None, fill=_fill(NAVY), border=BOX)
    sale_sum = "+".join(f"D{r}" for r in money_rows) if money_rows else "0"
    sh.put(4, f"={sale_sum}", font=_font(9, True, WHITE), fill=_fill(NAVY), border=BOX, fmt=MONEY0, align=RIGHT)
    for k, _ in cols:
        col = first_money_col + cols.index((k, _))
        formula = "=" + ("+".join(f"{letter[k]}{r}" for r in money_rows) if money_rows else "0")
        sh.put(col, formula, font=_font(9, True, WHITE), fill=_fill(NAVY), border=BOX, fmt=MONEY0, align=RIGHT)
    share = sh.row(13)
    sh.put(1, "% of contract price", font=_font(8, color=GREY, italic=True), align=RIGHT)
    for k, _ in cols:
        col = first_money_col + cols.index((k, _))
        sh.put(col, f"=IF($D${tot}=0,0,{letter[k]}{tot}/$D${tot})", font=_font(8, color=GREY, italic=True),
               fmt=PCT, align=RIGHT)
    _ = share
    sh.row(8)

    # ------------------------------------------------------ lower block --
    def _sum(keys: tuple[str, ...]) -> str:
        return "+".join(f"{letter[k]}{tot}" for k in keys)

    lower = s["lower"]
    unassigned_shown = any(k == "unassigned" for k, _ in cols)
    block: list[tuple[str, Any, bool]] = [
        ("TOTAL MATERIAL", f"={_sum(MATERIAL_COLUMNS)}", True),
        ("TOTAL LABOR", f"={_sum(LABOR_COLUMNS)}", True),
        ("TOTAL SUB LABOR", f"={_sum(SUB_LABOR_COLUMNS)}", True),
        ("TOTAL OTHER", f"={_sum(OTHER_COLUMNS)}", True),
        ("SALES TAX", f"={letter['tax']}{tot}", True),
    ]
    if unassigned_shown:
        block.append(("UNASSIGNED", f"={letter['unassigned']}{tot}", True))
    block.append(("MARGIN CONTINGENCY", _num(lower["margin_contingency"]), True))
    if Decimal(str(lower["difference"])) != 0:
        block.append(("ROUNDING", _num(lower["difference"]), True))
    first_block = sh.r + 1
    rows_in: list[int] = []
    for label, value, _counted in block:
        r = sh.row(15)
        rows_in.append(r)
        sh.put(2, label, font=_font(9, True))
        sh.put(4, value, font=_font(9), fmt=MONEY0, align=RIGHT, border=BOX)
        sh.put(5, f"=IF($D${tot}=0,0,D{r}/$D${tot})", font=_font(8, color=GREY), fmt=PCT, align=RIGHT)
    r = sh.row(16)
    sh.put(2, "ESTIMATED PROFIT", font=_font(10, True, NAVY))
    sh.put(4, f"=D{tot}-SUM(D{first_block}:D{rows_in[-1]})", font=_font(10, True, NAVY), fmt=MONEY0,
           align=RIGHT, border=BOX)
    sh.put(5, f"=IF($D${tot}=0,0,D{r}/$D${tot})", font=_font(8, color=GREY), fmt=PCT, align=RIGHT)
    r = sh.row(15)
    sh.put(2, "CONTRACT PRICE", font=_font(10, True))
    sh.put(4, f"=D{tot}", font=_font(10, True), fmt=MONEY0, align=RIGHT, border=BOX)
    r = sh.row(14)
    pct = float(Decimal(str(lower["labor_insurance_pct"])) * 100)
    sh.put(2, f"LABOR INSURANCE ({pct:g}% of sub labor, labor and supervision; not deducted)",
           font=_font(8, color=GREY, italic=True))
    sh.put(4, f"=({letter['sub_labor']}{tot}+{letter['labor']}{tot}+{letter['supervision']}{tot})*{lower['labor_insurance_pct']}",
           font=_font(8, color=GREY, italic=True), fmt=MONEY0, align=RIGHT)
    sh.row(10)

    # --------------------------------------------------------- cost codes --
    sh.row(30)
    for i, text in enumerate(("ITEM", "CODE", "NAME", "SUMMARY"), start=1):
        sh.put(i, text, font=_font(8, True, WHITE), fill=_fill(NAVY), align=CENTER if i > 1 else LEFT, border=BOX)
    for j, sec in enumerate(sections):
        sh.put(5 + j, sec["name"], font=_font(8, True, WHITE), fill=_fill(NAVY), align=CENTER, border=BOX)
    ws.column_dimensions["C"].width = max(ws.column_dimensions["C"].width or 0, 26)
    sec_first, sec_last = get_column_letter(5), get_column_letter(4 + max(1, len(sections)))

    code_rows: dict[str, int] = {}
    subtotal_rows: list[tuple[int, str]] = []
    last_category = None
    for cc in s["codes"]:
        r = sh.row(14)
        code_rows[cc["code"]] = r
        if cc["category"] != last_category:
            sh.put(1, cc["category_label"], font=_font(8, True, GREY))
            last_category = cc["category"]
        sub = bool(cc["is_subtotal"])
        f = _font(9, sub, GREY if sub else DARK, italic=sub)
        sh.put(2, cc["code"], font=f, border=BOX)
        sh.put(3, cc["name"], font=f, border=BOX)
        if sub:
            subtotal_rows.append((r, cc["category"]))   # its block is written below it; filled in after
        else:
            for j, sec in enumerate(sections):
                sh.put(5 + j, _num(sec["codes"].get(cc["code"], 0)), font=f, border=BOX, fmt=MONEY0, align=RIGHT)
        sh.put(4, f"=SUM({sec_first}{r}:{sec_last}{r})" if sections else 0, font=_font(9, True, GREY if sub else DARK, italic=sub),
               border=BOX, fmt=MONEY0, align=RIGHT)
    # A subtotal (000040, 000060) is the sum of its block, live, and is not added again.
    for r, category in subtotal_rows:
        block_codes = [o["code"] for o in s["codes"] if o["category"] == category and not o["is_subtotal"]]
        f = _font(9, True, GREY, italic=True)
        for j in range(max(1, len(sections))):
            col = get_column_letter(5 + j)
            parts = "+".join(f"{col}{code_rows[c]}" for c in block_codes if c in code_rows)
            ws.cell(r, 5 + j).value = f"={parts}" if parts and sections else 0
            ws.cell(r, 5 + j).font = f
            ws.cell(r, 5 + j).border = BOX
            ws.cell(r, 5 + j).number_format = MONEY0
            ws.cell(r, 5 + j).alignment = RIGHT

    # The rows the chart has no code for, then the tie-out.
    extra: list[tuple[str, str, str]] = [("", "Sales tax", "tax")]
    if unassigned_shown:
        extra.append(("", "Unassigned", "unassigned"))
    extra_rows: dict[str, int] = {}
    for item, name, key in extra:
        r = sh.row(14)
        extra_rows[key] = r
        sh.put(1, item, font=_font(8, True, GREY))
        sh.put(2, "", border=BOX)
        sh.put(3, name, font=_font(9), border=BOX)
        for j, sec in enumerate(sections):
            sh.put(5 + j, _num(sec["columns"].get(key, 0)), font=_font(9), border=BOX, fmt=MONEY0, align=RIGHT)
        sh.put(4, f"=SUM({sec_first}{r}:{sec_last}{r})" if sections else 0, font=_font(9, True), border=BOX,
               fmt=MONEY0, align=RIGHT)
    if any(Decimal(str(sec["difference"])) != 0 for sec in sections):
        r = sh.row(14)
        extra_rows["difference"] = r
        sh.put(2, "", border=BOX)
        sh.put(3, "Rounding (stored cost less the lines)", font=_font(9, color=GREY, italic=True), border=BOX)
        for j, sec in enumerate(sections):
            sh.put(5 + j, _num(sec["difference"]), font=_font(9, color=GREY, italic=True), border=BOX, fmt=MONEY2, align=RIGHT)
        sh.put(4, f"=SUM({sec_first}{r}:{sec_last}{r})" if sections else 0, font=_font(9, color=GREY, italic=True),
               border=BOX, fmt=MONEY2, align=RIGHT)

    counted = [code_rows[c["code"]] for c in s["codes"] if not c["is_subtotal"] and c["code"] != "000091"]
    counted += list(extra_rows.values())
    r = sh.row(16)
    sh.put(2, "", fill=_fill(LIGHT), border=BOX)
    sh.put(3, "TOTAL COST", font=_font(9, True), fill=_fill(LIGHT), border=BOX)
    for j, sec in enumerate(sections):
        col = get_column_letter(5 + j)
        sh.put(5 + j, "=" + "+".join(f"{col}{x}" for x in counted), font=_font(9, True), fill=_fill(LIGHT),
               border=BOX, fmt=MONEY0, align=RIGHT)
    sh.put(4, f"=SUM({sec_first}{r}:{sec_last}{r})" if sections else 0, font=_font(9, True), fill=_fill(LIGHT),
           border=BOX, fmt=MONEY0, align=RIGHT)
    cost_row = r
    r = sh.row(16)
    sh.put(2, "", border=BOX)
    sh.put(3, "SALE", font=_font(9, True), border=BOX)
    for j, sec in enumerate(sections):
        sh.put(5 + j, _num(sec["sale"]), font=_font(9, True), border=BOX, fmt=MONEY0, align=RIGHT)
    sh.put(4, f"=SUM({sec_first}{r}:{sec_last}{r})" if sections else 0, font=_font(9, True), border=BOX,
           fmt=MONEY0, align=RIGHT)
    sale_row = r
    r = sh.row(16)
    sh.put(2, "", border=BOX)
    sh.put(3, "ESTIMATED PROFIT", font=_font(9, True, NAVY), border=BOX)
    for j, sec in enumerate(sections):
        col = get_column_letter(5 + j)
        sh.put(5 + j, f"={col}{sale_row}-{col}{cost_row}-{col}{code_rows['000091']}", font=_font(9, True, NAVY),
               border=BOX, fmt=MONEY0, align=RIGHT)
    sh.put(4, f"=SUM({sec_first}{r}:{sec_last}{r})" if sections else 0, font=_font(9, True, NAVY), border=BOX,
           fmt=MONEY0, align=RIGHT)

    # The page: landscape, one page wide.
    ws.freeze_panes = "B1"
    ws.print_area = f"A1:{get_column_letter(max(last_col, 4 + len(sections)))}{sh.r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "CONFIDENTIAL"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
