"""
An estimate workbook, read onto a job (2026-09-10).

Chad, 2026-09-10, with the Lakeside Townhomes workbook (26-051, rev 14):
"lets import this estimate" — and, on the proposal, "full import, build it".

## What this reads

The office's estimate workbook: one tab per assembly, a takeoff grid under
a header row, and under the grid the tab's own money — the mixes it buys
and their prices, the steel, the PT, the sand, the waste factors, the labor
block (a rate per line and a Y/N for subbing it), the supervision block
(days at a day rate), the equipment block (days at a day rate), the
contract services, and the PRICE row with the section's sale and its
COST + markup.

Tabs are found by name (a mono slab, a paving, the sidewalks, the grade
beams, the footings, a slab on grade, the miscellaneous) and their columns
by HEADER TEXT, not by letter — the Lakeside template has the headers a row
lower than the LBJ one and a building-type quantity column the older tab
lacks, and the next template will move something else. A tab with no
priced rows is skipped and said so.

## What it writes

Through the app's own services, never around them: the project, the
estimate (which pulls the master price list), the job's price sheet at the
workbook's prices, one section per tab with the tab's waste and markup and
labor subbed as the tab says, the takeoff rows (a mono slab's beam schedule
and its per-pour beam, exposed-beam and drop feet included; a garden-style
pour at its building count), the three line sets built the way opening the
section builds them, then the tab's rates onto the section — the labor
rates as section rates, the supervision and equipment DAYS typed onto
their lines, the contract rates as section rates — and the totals rolled
up. Quantities the app derives (square feet, pump yards, saw-cut feet, the
tons of steel) are left derived: the app keeps its rule and the tie-out
names the difference.

The tie-out is each section's app sale against the tab's SALE cell, the
way the LBJ reconciliation was done; the report carries both and the gap.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from sqlalchemy import select, text
from sqlalchemy.orm import Session

_Q2 = Decimal("0.01")
_Q4 = Decimal("0.0001")
_ZERO = Decimal("0")


# ------------------------------------------------------------ the cells ----


def norm(v: Any) -> str:
    """Upper, one space between words, no stray punctuation at the ends."""
    if v is None:
        return ""
    s = str(v).replace("’", "'").upper()
    s = re.sub(r"\s+", " ", s).strip()
    return s.strip(" :.")


def num(v: Any) -> Decimal | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        d = Decimal(str(v))
        return d if d.is_finite() else None
    try:
        return Decimal(str(v).replace(",", "").replace("$", "").strip())
    except Exception:  # noqa: BLE001 — a label in a number's cell is not a number
        return None


def pos(v: Any) -> Decimal | None:
    d = num(v)
    return d if d is not None and d > 0 else None


def yes(v: Any) -> bool:
    return norm(v) in ("Y", "YES", "TRUE", "X")


def ival(v: Any) -> int | None:
    d = num(v)
    return int(d) if d is not None and d == d.to_integral_value() and d > 0 else None


def header_row(ws, *must: str, rows=range(1, 16)) -> tuple[int, dict[str, int]] | None:
    """The first row carrying every one of `must` as a cell label (a prefix counts)."""
    for r in rows:
        cols: dict[str, int] = {}
        for c in range(1, min(ws.max_column, 80) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                cols.setdefault(norm(v), c)
        if all(any(k == norm(m) or k.startswith(norm(m)) for k in cols) for m in must):
            return r, cols
    return None


def col(cols: dict[str, int], *names: str) -> int | None:
    for n in names:
        n = norm(n)
        if n in cols:
            return cols[n]
    for n in names:
        n = norm(n)
        for k, c in cols.items():
            if k.startswith(n):
                return c
    return None


def cell(ws, r: int, c: int | None) -> Any:
    return ws.cell(r, c).value if c else None


STOP_WORDS = ("TOTAL SQ FT", "TOTAL SF", "TOTAL LINEAR", "TOTAL LF", "TOTAL FF", "UNIT SALE")


def data_rows(ws, start: int, key_col: int, limit: int = 60):
    """Rows with a positive number in the key column, until the tab's totals."""
    for r in range(start, start + limit):
        labels = [norm(ws.cell(r, c).value) for c in range(1, 10)]
        if any(lb.startswith(w) for lb in labels for w in STOP_WORDS):
            break
        if pos(ws.cell(r, key_col).value) is None:
            continue
        yield r


def find_label(ws, needle: str, *, cols=(1, 2, 3), rows=None) -> tuple[int, int] | None:
    """The first cell in the given columns whose text starts with `needle`."""
    rows = rows or range(1, ws.max_row + 1)
    n = norm(needle)
    for r in rows:
        for c in cols:
            v = ws.cell(r, c).value
            if isinstance(v, str) and norm(v).startswith(n):
                return r, c
    return None


# ------------------------------------------------------------- the spec ----


@dataclass
class TabSpec:
    tab: str
    kind: str
    name: str
    unit: str
    rows: list[dict[str, Any]] = field(default_factory=list)
    beam_types: list[dict[str, Any]] = field(default_factory=list)      # mono slab schedule
    usages: dict[int, list[tuple[str, int, Decimal]]] = field(default_factory=dict)  # row -> (kind, type#, LF)
    mixes: dict[int, tuple[str, Decimal | None]] = field(default_factory=dict)
    prices: dict[str, Decimal] = field(default_factory=dict)            # steel_lb, mesh_sf, pt_sf, sand_cy, rock_cy
    waste: dict[str, Decimal] = field(default_factory=dict)             # concrete, rebar, sand
    switches: dict[str, bool] = field(default_factory=dict)             # carton_forms, durrock_retainer: the tab's Y/N
    labor: list[dict[str, Any]] = field(default_factory=list)
    supervision: list[dict[str, Any]] = field(default_factory=list)
    equipment: list[dict[str, Any]] = field(default_factory=list)
    contract: list[dict[str, Any]] = field(default_factory=list)
    margin: Decimal | None = None
    sale: Decimal | None = None
    unit_sale: Decimal | None = None
    quantity: Decimal | None = None
    tax_exempt: bool | None = None
    supplier: str | None = None
    notes: list[str] = field(default_factory=list)

    @property
    def total_cost(self) -> Decimal | None:
        if self.sale is None:
            return None
        m = self.margin or _ZERO
        return (self.sale / (1 + m)).quantize(_Q2)


@dataclass
class JobSpec:
    path: str
    project: dict[str, Any]
    estimate: dict[str, Any]
    tabs: list[TabSpec] = field(default_factory=list)
    skipped: list[tuple[str, str]] = field(default_factory=list)
    lumber: dict[str, Decimal] = field(default_factory=dict)            # Pricing tab, by name
    day_rates: dict[str, Decimal] = field(default_factory=dict)         # the tabs' supervision day rates


# Tab name -> section kind. Order matters: "Gd Beams or Cont Footings" is
# footings, "Garage Walls" is walls, "Site Walls" is walls, "Footings" alone
# is spot footings.
KIND_BY_NAME: tuple[tuple[str, str], ...] = (
    ("CONT FOOTING", "cont_footings"),
    ("GD BEAM", "grade_beams"),
    ("GRADE BEAM", "grade_beams"),
    ("MONO SLAB", "mono_slab"),
    ("PT SLAB", "mono_slab"),
    ("SLAB ON DECK", "slab_on_deck"),
    ("SLAB ON GRADE", "slabs"),
    ("SLABS", "slabs"),
    ("PAVING", "paving"),
    ("SIDEWALK", "sidewalk"),
    ("WALL", "walls_footings"),
    ("FOOTING", "spot_footings"),
    ("PIER", "piers"),
    ("COLUMN", "columns"),
    ("CIP EL", "cip_deck"),
    ("PANEL", "panels"),
    ("MISC", "miscellaneous"),
)
NOT_TABS = ("SUMMARY", "PRICING", "INFORMATION", "CHECK LIST", "SUB LABOR", "TIMELINE", "MATERIALS",
            "CONCRETE YARDAGE", "CONCRETE BIDS", "PROPOSAL", "CONTINGENCY")
UNIT_BY_KIND = {
    "mono_slab": "SF", "paving": "SF", "sidewalk": "SF", "slabs": "SF", "grade_beams": "LF",
    "cont_footings": "LF", "spot_footings": "EA", "walls_footings": "FF", "miscellaneous": "LS",
}
READERS_BUILT = {"mono_slab", "paving", "sidewalk", "slabs", "grade_beams", "cont_footings", "spot_footings", "miscellaneous"}


def kind_of(tab_name: str) -> str | None:
    n = norm(tab_name)
    if any(n.startswith(x) or x in n for x in NOT_TABS):
        return None
    for needle, kind in KIND_BY_NAME:
        if needle in n:
            return kind
    return None


# --------------------------------------------------------- the blocks ----


def money_block(ws, t: TabSpec) -> None:
    """The mixes and their prices, the steel, mesh, PT, sand and rock, and the waste factors."""
    hit = find_label(ws, "MIX #", rows=range(20, ws.max_row + 1))
    if hit is None:
        t.notes.append("no material block")
        return
    r0, L = hit
    for r in range(r0 + 1, r0 + 20):
        for cc in (L, L + 1):
            lab = norm(cell(ws, r, cc))
            if "CARTON" in lab or "RETAINER" in lab or "DURROCK" in lab:
                flag = next((norm(cell(ws, r, c)) for c in range(cc + 1, cc + 6) if norm(cell(ws, r, c)) in ("Y", "N")), None)
                if flag is not None:
                    t.switches["carton_forms" if "CARTON" in lab else "durrock_retainer"] = flag == "Y"
    for r in range(r0 + 1, r0 + 16):
        n = ival(cell(ws, r, L))
        name = cell(ws, r, L + 1)
        if n is not None and isinstance(name, str) and name.strip():
            t.mixes[n] = (name.strip(), num(cell(ws, r, L + 5)))
            if "concrete" not in t.waste:
                w = num(cell(ws, r, L + 9))
                if w is not None and norm(cell(ws, r, L + 7)).startswith("CONCRETE WASTAGE"):
                    t.waste["concrete"] = w
            continue
        for cc in (L, L + 1):
            lab = norm(cell(ws, r, cc))
            if not lab:
                continue
            price = num(cell(ws, r, L + 5))
            waste = num(cell(ws, r, L + 9)) if "WASTAGE" in norm(cell(ws, r, L + 7)) else None
            if "STEEL" in lab and price is not None:
                t.prices["steel_lb"] = price
                if waste is not None:
                    t.waste["rebar"] = waste
            elif "WIRE MESH" in lab and price is not None:
                t.prices["mesh_sf"] = price
            elif lab.startswith("POST") and price is not None:
                t.prices["pt_sf"] = price
            elif "SAND" in lab and price is not None:
                t.prices["sand_cy"] = price
                if waste is not None:
                    t.waste["sand"] = waste
            elif "ROCK" in lab and price is not None:
                t.prices["rock_cy"] = price


def _label_at(ws, r: int, L: int) -> str | None:
    v = cell(ws, r, L)
    return v.strip() if isinstance(v, str) and v.strip() else None


def labor_blocks(ws, t: TabSpec) -> None:
    """The labor, supervision, equipment and contract blocks under the grid, by their offsets from the label column."""
    hit = find_label(ws, "LABOR:", rows=range(20, ws.max_row + 1))
    if hit is None:
        t.notes.append("no labor block")
        return
    r0, L = hit
    sub_col = house_col = None
    for c in range(L + 1, min(ws.max_column, L + 24) + 1):
        lab = norm(cell(ws, r0, c))
        if lab == "SUB LABOR":
            sub_col = c
        elif lab == "LABOR":
            house_col = c
    sub_col = sub_col or L + 13
    house_col = house_col or L + 10

    def money(r: int, c: int | None) -> Decimal | None:
        """The first number under or beside a header — the tabs merge the cost cells."""
        if c is None:
            return None
        for cc in (c, c + 1, c + 2):
            v = num(cell(ws, r, cc))
            if v is not None:
                return v
        return None

    def cost(r: int) -> Decimal | None:
        subbed, house = money(r, sub_col), money(r, house_col)
        return subbed if subbed else house

    r = r0 + 1
    end = min(ws.max_row, r0 + 80)
    # labor
    while r <= end:
        lab = _label_at(ws, r, L)
        if lab and norm(lab).startswith("COST PER"):
            break
        if lab:
            t.labor.append({
                "label": lab, "sub": yes(cell(ws, r, L + 2)), "rate": num(cell(ws, r, L + 3)),
                "unit": cell(ws, r, L + 5), "qty": num(cell(ws, r, L + 7)), "cost": cost(r),
            })
        r += 1
    # supervision
    while r <= end and not norm(_label_at(ws, r, L) or "").startswith("SUPERVISION"):
        r += 1
    r += 1
    while r <= end:
        lab = _label_at(ws, r, L)
        if lab and norm(lab).startswith("EQUIPMENT"):
            break
        if lab:
            t.supervision.append({"label": lab, "days": num(cell(ws, r, L + 3)), "rate": num(cell(ws, r, L + 5)),
                                  "cost": money(r, sub_col)})
        r += 1
    r += 1
    # equipment
    while r <= end:
        lab = _label_at(ws, r, L)
        if lab and norm(lab).startswith("CONTRACT"):
            break
        if lab:
            t.equipment.append({"label": lab, "days": num(cell(ws, r, L + 3)), "rate": num(cell(ws, r, L + 5)),
                                "fuel": num(cell(ws, r, L + 8)), "cost": money(r, sub_col)})
        r += 1
    r += 1
    # contract services, until the tab's price row
    while r <= end:
        lab = _label_at(ws, r, L)
        if lab and (norm(lab).startswith("PRICE") or norm(lab).startswith("MARGIN") or norm(lab).startswith("UNIT SALE")):
            break
        if lab:
            n = norm(lab)
            if n.startswith("OUT OF TOWN") or n.startswith("MISC"):
                item = {"label": lab, "qty": num(cell(ws, r, L + 3)), "rate": num(cell(ws, r, L + 5))}
            else:
                item = {"label": lab, "rate": num(cell(ws, r, L + 3)), "unit": cell(ws, r, L + 4),
                        "qty": num(cell(ws, r, L + 5))}
            item["cost"] = money(r, sub_col)
            item["enabled"] = True
            # "Pump Paving?  N" beside the pump line turns it off.
            for cc in range(L + 8, L + 14):
                if "PUMP" in norm(cell(ws, r, cc)) and "?" in str(cell(ws, r, cc)):
                    item["enabled"] = yes(cell(ws, r, cc + 1))
            t.contract.append(item)
        r += 1


def price_row(ws, t: TabSpec) -> None:
    """The PRICE: row — the unit sale, the sale, the COST + markup — and the tax flag."""
    hit = find_label(ws, "PRICE:", rows=range(20, ws.max_row + 1))
    if hit is None:
        t.notes.append("no PRICE row")
        return
    r, L = hit
    t.unit_sale = num(cell(ws, r, L + 1))
    t.sale = num(cell(ws, r, L + 3))
    for c in range(1, min(ws.max_column, 40) + 1):
        if norm(cell(ws, r - 1, c)) == "COST +":
            t.margin = num(cell(ws, r, c))
    for rr in range(r - 2, r + 2):
        for c in range(1, min(ws.max_column, 60) + 1):
            if norm(cell(ws, rr, c)) == "EXEMPT":
                t.tax_exempt = yes(cell(ws, rr, c + 1))
    sup = find_label(ws, "CONCRETE SUPPLIER", cols=range(1, 12), rows=range(1, 12))
    if sup is not None:
        rr, cc = sup
        for c in range(cc + 1, cc + 8):
            v = cell(ws, rr, c)
            if isinstance(v, str) and v.strip():
                t.supplier = v.strip()
                break


# -------------------------------------------------------- the readers ----


def _name(ws, r: int, *cols: int | None) -> str:
    parts = []
    for c in cols:
        v = cell(ws, r, c)
        if v not in (None, ""):
            parts.append(str(v).strip())
    return " ".join(parts)


def read_mono_slab(ws, t: TabSpec) -> None:
    """A slab tab: the pours (a building type at its count), the beam schedule, and each pour's beam feet."""
    hdr = header_row(ws, "SQUARE FOOTAGE", "THICK INCH")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    h, cols = hdr
    c_type = col(cols, "TYPE", "PT SOG TYPE")
    c_bld = col(cols, "BLD #", "BLD")
    c_sf = col(cols, "SQUARE FOOTAGE", "SQ")
    c_thick = col(cols, "THICK INCH", "THICK INCHES")
    c_qty = col(cols, "QTY")
    c_cable = col(cols, "CABLE")
    c_mix = col(cols, "MIX DESIGN")
    c_sand = col(cols, "INCH OF SAND", "INCH OF", "SAND INCHE", "SAND INCHES")
    c_perim = col(cols, "PERM. EDGE", "PERM EDGE")
    c_reinf = col(cols, "REINFORCING", "RE")
    c_mesh = col(cols, "WIRE MESH")
    c_exp = col(cols, "EXP GB")
    c_drops = col(cols, "DROPS")
    c_gbs = col(cols, "GRADE BEAMS")
    c_add = col(cols, "LABOR ADD", "COST ADDER", "PAVING ADD")
    sub = h + 1
    pairs: list[int] = []
    if c_gbs:
        c = c_gbs
        while c < ws.max_column and norm(cell(ws, sub, c)) == "TYPE" and len(pairs) < 8:
            pairs.append(c)
            c += 2
    for r in data_rows(ws, sub + 1, c_sf):
        row = {
            "description": _name(ws, r, c_type, c_bld) or f"Pour {r}",
            "square_footage": num(cell(ws, r, c_sf)),
            "thickness_in": num(cell(ws, r, c_thick)) or Decimal("4"),
            "qty": ival(cell(ws, r, c_qty)) or 1,
            "post_tension": yes(cell(ws, r, c_cable)),
            "mix": ival(cell(ws, r, c_mix)),
            "sand_thickness_in": num(cell(ws, r, c_sand)),
            "perimeter_edge_lf": num(cell(ws, r, c_perim)),
            "slab_bar_size": ival(cell(ws, r, c_reinf)) if c_reinf else None,
            "slab_bar_spacing_in": num(cell(ws, r, c_reinf + 1)) if c_reinf else None,
            "wire_mesh": yes(cell(ws, r, c_mesh)) if c_mesh else False,
            "paving_add_per_sf": pos(cell(ws, r, c_add)) if c_add else None,
        }
        t.rows.append(row)
        uses: list[tuple[str, int, Decimal]] = []
        for c in pairs:
            n, lf = ival(cell(ws, r, c)), pos(cell(ws, r, c + 1))
            if n and lf:
                uses.append(("grade_beam", n, lf))
        if c_exp:
            n, lf = ival(cell(ws, r, c_exp)), pos(cell(ws, r, c_exp + 1))
            if n and lf:
                uses.append(("exposed", n, lf))
        if c_drops:
            n, lf = ival(cell(ws, r, c_drops)), pos(cell(ws, r, c_drops + 1))
            if n and lf:
                uses.append(("drop", n, lf))
        if uses:
            t.usages[len(t.rows) - 1] = uses
    # The beam schedule: GB # | width | height | top n, size | bot n, size | mid n, size | stirrup size, spacing | L bars size, spacing, length
    hit = find_label(ws, "GB #", rows=range(h + 2, min(ws.max_row, h + 120)))
    if hit is not None:
        r0, L = hit
        for r in range(r0 + 2, r0 + 14):
            n = ival(cell(ws, r, L))
            w, hgt = pos(cell(ws, r, L + 1)), pos(cell(ws, r, L + 2))
            if n is None or w is None or hgt is None:
                continue
            t.beam_types.append({
                "n": n, "width_in": w, "height_in": hgt,
                "top_bars_count": ival(cell(ws, r, L + 3)) or 0, "top_bars_size": ival(cell(ws, r, L + 4)),
                "bottom_bars_count": ival(cell(ws, r, L + 5)) or 0, "bottom_bars_size": ival(cell(ws, r, L + 6)),
                "mid_bars_count": ival(cell(ws, r, L + 7)) or 0, "mid_bars_size": ival(cell(ws, r, L + 8)),
                "stirrup_size": ival(cell(ws, r, L + 9)), "stirrup_spacing_in": num(cell(ws, r, L + 10)),
                "l_bars_size": ival(cell(ws, r, L + 11)), "l_bars_spacing_in": num(cell(ws, r, L + 12)),
                "l_bars_length_ft": num(cell(ws, r, L + 13)),
            })
    t.quantity = sum((row["square_footage"] * row["qty"] for row in t.rows), _ZERO)


def read_slabs(ws, t: TabSpec) -> None:
    """A slab on grade tab: conventional slabs with a bar mat, a thickened edge, cartons."""
    hdr = header_row(ws, "SQUARE FOOTAGE", "THICK INCHES")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    h, cols = hdr
    c_type, c_sf = col(cols, "TYPE"), col(cols, "SQUARE FOOTAGE")
    c_thick, c_perim, c_mix = col(cols, "THICK INCHES"), col(cols, "PERM. EDGE", "PERM EDGE"), col(cols, "MIX DESIGN")
    c_edge, c_sand = col(cols, "LN FT THICK EDGE"), col(cols, "SAND INCHE", "SAND INCHES")
    c_top, c_bot, c_mesh = col(cols, "TOP REINFORCING"), col(cols, "BOTTOM REINFORCING"), col(cols, "WIRE MESH")
    c_carton = col(cols, "CARTON")
    c_add = col(cols, "LABOR ADD", "COST ADDER", "PAVING ADD")
    for r in data_rows(ws, h + 2, c_sf):
        size = ival(cell(ws, r, c_bot)) if c_bot else None
        spacing = num(cell(ws, r, c_bot + 1)) if c_bot else None
        if size is None and c_top:
            size, spacing = ival(cell(ws, r, c_top)), num(cell(ws, r, c_top + 1))
        t.rows.append({
            "description": _name(ws, r, c_type) or f"Slab {r}",
            "square_footage": num(cell(ws, r, c_sf)),
            "thickness_in": num(cell(ws, r, c_thick)) or Decimal("4"),
            "qty": 1, "post_tension": False,
            "mix": ival(cell(ws, r, c_mix)),
            "sand_thickness_in": num(cell(ws, r, c_sand)),
            "perimeter_edge_lf": num(cell(ws, r, c_perim)),
            "thick_edge_lf": num(cell(ws, r, c_edge)),
            "slab_bar_size": size, "slab_bar_spacing_in": spacing,
            "wire_mesh": yes(cell(ws, r, c_mesh)) if c_mesh else False,
            "carton": yes(cell(ws, r, c_carton)) if c_carton else False,
            "paving_add_per_sf": pos(cell(ws, r, c_add)) if c_add else None,
        })
    t.quantity = sum((row["square_footage"] for row in t.rows), _ZERO)


def read_paving(ws, t: TabSpec) -> None:
    hdr = header_row(ws, "SQUARE FOOTAGE", "LF CURBS")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    h, cols = hdr
    c_type, c_sf, c_thick = col(cols, "PAVING TYPE", "TYPE"), col(cols, "SQUARE FOOTAGE"), col(cols, "THICK INCHES")
    c_slip, c_traffic = col(cols, "SLIP FORM"), col(cols, "TRAFFIC CONTROL")
    c_demo, c_curb, c_mix = col(cols, "LN FT DEMO"), col(cols, "LF CURBS"), col(cols, "MIX DESIGN")
    c_edge, c_sand = col(cols, "LN FT THICK EDGE"), col(cols, "SAND INCHE", "SAND INCHES")
    c_steel, c_gauge, c_add = col(cols, "STEEL REINFORCING", "REINFORCING"), col(cols, "MESH GAUGE"), col(cols, "PAVING ADD")
    for r in data_rows(ws, h + 2, c_sf):
        t.rows.append({
            "description": _name(ws, r, c_type) or f"Area {r}",
            "square_footage": num(cell(ws, r, c_sf)),
            "thickness_in": num(cell(ws, r, c_thick)) or Decimal("5"),
            "qty": 1, "post_tension": False,
            "slip_form": yes(cell(ws, r, c_slip)) if c_slip else False,
            "traffic_control": yes(cell(ws, r, c_traffic)) if c_traffic else False,
            "demo_lf": num(cell(ws, r, c_demo)),
            "curb_lf": num(cell(ws, r, c_curb)),
            "mix": ival(cell(ws, r, c_mix)),
            "thick_edge_lf": num(cell(ws, r, c_edge)),
            "sand_thickness_in": num(cell(ws, r, c_sand)),
            "slab_bar_size": ival(cell(ws, r, c_steel)) if c_steel else None,
            "slab_bar_spacing_in": num(cell(ws, r, c_steel + 1)) if c_steel else None,
            "mesh_gauge": ival(cell(ws, r, c_gauge)) if c_gauge else None,
            "paving_add_per_sf": num(cell(ws, r, c_add)) if c_add else None,
        })
    t.quantity = sum((row["square_footage"] for row in t.rows), _ZERO)


def read_sidewalks(ws, t: TabSpec) -> None:
    hdr = header_row(ws, "SQUARE FOOTAGE", "STAIR TREADS")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    h, cols = hdr
    c_type, c_sf, c_thick = col(cols, "SIDEWALK TYPE", "TYPE"), col(cols, "SQUARE FOOTAGE"), col(cols, "THICK INCHES")
    c_sand, c_mix, c_edge = col(cols, "SAND INCHE", "SAND INCHES"), col(cols, "MIX DESIGN"), col(cols, "LN FT THICK EDGE")
    c_stamp, c_color, c_acid = col(cols, "STAMPED"), col(cols, "INTRAGAL COLOR", "INTEGRAL COLOR"), col(cols, "ACID")
    c_traffic, c_stairs, c_reinf = col(cols, "TRAFFIC CONTROL"), col(cols, "STAIR TREADS"), col(cols, "REINFORCING")
    c_add = col(cols, "COST ADDER", "LABOR ADD", "PAVING ADD")
    for r in data_rows(ws, h + 2, c_sf):
        t.rows.append({
            "description": _name(ws, r, c_type) or f"Walk {r}",
            "square_footage": num(cell(ws, r, c_sf)),
            "thickness_in": num(cell(ws, r, c_thick)) or Decimal("4"),
            "qty": 1, "post_tension": False,
            "sand_thickness_in": num(cell(ws, r, c_sand)),
            "mix": ival(cell(ws, r, c_mix)),
            "thick_edge_lf": num(cell(ws, r, c_edge)),
            "stamped": yes(cell(ws, r, c_stamp)) if c_stamp else False,
            "integral_color": yes(cell(ws, r, c_color)) if c_color else False,
            "acid_etch": yes(cell(ws, r, c_acid)) if c_acid else False,
            "traffic_control": yes(cell(ws, r, c_traffic)) if c_traffic else False,
            "stair_tread_lf": num(cell(ws, r, c_stairs)) if c_stairs else None,
            "stair_tread_rise_in": num(cell(ws, r, c_stairs + 1)) if c_stairs else None,
            "stair_tread_run_in": num(cell(ws, r, c_stairs + 2)) if c_stairs else None,
            "slab_bar_size": ival(cell(ws, r, c_reinf)) if c_reinf else None,
            "slab_bar_spacing_in": num(cell(ws, r, c_reinf + 1)) if c_reinf else None,
            "mesh_gauge": ival(cell(ws, r, c_reinf + 2)) if c_reinf else None,
            "paving_add_per_sf": pos(cell(ws, r, c_add)) if c_add else None,
        })
    t.quantity = sum((row["square_footage"] for row in t.rows), _ZERO)


def read_beams(ws, t: TabSpec) -> None:
    """The beam tab: a type per row with its length — a beam run (sql/073)."""
    hdr = header_row(ws, "GB TYPE", "LENGTH FT")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    s, sub = hdr
    g = header_row(ws, "TOP BARS", rows=range(max(1, s - 2), s))
    grp = g[1] if g else {}
    c_type, c_len = col(sub, "GB TYPE"), col(sub, "LENGTH FT")
    c_w, c_h = col(sub, "WIDTH IN"), col(sub, "HEIGHT IN")
    c_mix = col(grp, "MIX DESIGN") or col(sub, "MIX DESIGN")
    c_top, c_bot, c_mid = col(grp, "TOP BARS"), col(grp, "BOTTOM BARS"), col(grp, "MID BARS")
    c_stir, c_l, c_pil = col(grp, "STIRRUPS"), col(grp, '"L" BARS', "L BARS"), col(grp, "PILASTERS")
    for r in data_rows(ws, s + 1, c_len):
        t.rows.append({
            "label": _name(ws, r, c_type) or f"Beam {r}",
            "mix": ival(cell(ws, r, c_mix)),
            "length_ft": num(cell(ws, r, c_len)),
            "width_in": num(cell(ws, r, c_w)) or _ZERO, "height_in": num(cell(ws, r, c_h)) or _ZERO,
            "top_bars_count": ival(cell(ws, r, c_top)) or 0 if c_top else 0,
            "top_bars_size": ival(cell(ws, r, c_top + 1)) if c_top else None,
            "bottom_bars_count": ival(cell(ws, r, c_bot)) or 0 if c_bot else 0,
            "bottom_bars_size": ival(cell(ws, r, c_bot + 1)) if c_bot else None,
            "mid_bars_count": ival(cell(ws, r, c_mid)) or 0 if c_mid else 0,
            "mid_bars_size": ival(cell(ws, r, c_mid + 1)) if c_mid else None,
            "stirrup_size": ival(cell(ws, r, c_stir)) if c_stir else None,
            "stirrup_spacing_in": num(cell(ws, r, c_stir + 1)) if c_stir else None,
            "l_bars_size": ival(cell(ws, r, c_l)) if c_l else None,
            "l_bars_spacing_in": num(cell(ws, r, c_l + 1)) if c_l else None,
            "l_bars_length_ft": num(cell(ws, r, c_l + 2)) if c_l else None,
            "pilaster_count": ival(cell(ws, r, c_pil)) or 0 if c_pil else 0,
            "pilaster_length_in": num(cell(ws, r, c_pil + 1)) if c_pil else None,
            "pilaster_width_in": num(cell(ws, r, c_pil + 2)) if c_pil else None,
        })
    t.quantity = sum((row["length_ft"] for row in t.rows), _ZERO)


def read_footings(ws, t: TabSpec) -> None:
    """The footings tab: a type at a count, each so long and wide, so thick, two mats — a spot footing (sql/072)."""
    hdr = header_row(ws, "QTY", "LENGTH FT", "FEET WIDE")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    s, sub = hdr
    g = header_row(ws, "TOP MAT", rows=range(max(1, s - 2), s))
    grp = g[1] if g else {}
    c_type, c_qty, c_len = col(sub, "TYPE"), col(sub, "QTY"), col(sub, "LENGTH FT")
    c_wide, c_thick = col(sub, "FEET WIDE"), col(sub, "INCHES THICK")
    c_mix = col(grp, "MIX DESIGN")
    c_top, c_bot, c_pil = col(grp, "TOP MAT"), col(grp, "BOTTOM MAT"), col(grp, "PHILASTERS", "PILASTERS")
    for r in data_rows(ws, s + 1, c_len):
        count = ival(cell(ws, r, c_qty))
        if not count:
            continue
        wide = num(cell(ws, r, c_wide)) or _ZERO
        t.rows.append({
            "label": _name(ws, r, c_type) or f"Footing {r}",
            "mix": ival(cell(ws, r, c_mix)),
            "footing_count": count,
            "footing_each_ft": num(cell(ws, r, c_len)),
            "ftg_width_in": (wide * 12).quantize(Decimal("0.001")),
            "ftg_thick_in": num(cell(ws, r, c_thick)) or _ZERO,
            "ftg_top_spacing_in": num(cell(ws, r, c_top)) if c_top else None,
            "ftg_top_size": ival(cell(ws, r, c_top + 1)) if c_top else None,
            "ftg_bot_spacing_in": num(cell(ws, r, c_bot)) if c_bot else None,
            "ftg_bot_size": ival(cell(ws, r, c_bot + 1)) if c_bot else None,
            "pilaster_count": ival(cell(ws, r, c_pil)) or 0 if c_pil else 0,
        })
    t.quantity = Decimal(sum(row["footing_count"] for row in t.rows))


def read_misc(ws, t: TabSpec) -> None:
    """The miscellaneous tab: named items at a count, a typed sale and labor, a shape."""
    hdr = header_row(ws, "UNIT SALE", "LABOR")
    if hdr is None:
        t.notes.append("no takeoff header")
        return
    h, cols = hdr
    c_type = col(cols, "TYPE") or 2
    c_sub, c_qty, c_sale, c_labor = col(cols, "SUB LABOR"), col(cols, "QTY"), col(cols, "UNIT SALE"), col(cols, "LABOR COST")
    c_dia, c_depth, c_yds, c_steel = col(cols, "DIA"), col(cols, "DEPTH"), col(cols, "YDS CONC"), col(cols, "STEEL WT")
    # The tab's header names its site mix: "SITE CONCRETE MIX MIX DESIGN:  6  ...  3000 PSI Sidewalk and Hardscape  155".
    site_mix = None
    for r in range(max(1, h - 3), h):
        for c in range(1, min(ws.max_column, 40)):
            if "MIX DESIGN" in norm(cell(ws, r, c)):
                n = name = price = None
                for cc in range(c + 1, c + 12):
                    v = cell(ws, r, cc)
                    if n is None and ival(v) is not None and num(v) < 20:
                        n = ival(v)
                    elif name is None and isinstance(v, str) and v.strip() and "MIX" not in norm(v):
                        name = v.strip()
                    elif price is None and num(v) is not None and num(v) >= 20:
                        price = num(v)
                if n is not None and name:
                    t.mixes.setdefault(n, (name, price))
                    site_mix = site_mix or n
    for r in range(h + 1, h + 40):
        qty = pos(cell(ws, r, c_qty))
        name = None
        for c in range(1, (c_type or 2) + 1):
            v = cell(ws, r, c)
            if isinstance(v, str) and v.strip() and not norm(v).startswith("TYPE"):
                name = v.strip()
                break
        if not qty or not name:
            continue
        t.rows.append({
            "description": name, "qty": qty, "subcontracted": yes(cell(ws, r, c_sub)) if c_sub else True,
            "unit_sale": num(cell(ws, r, c_sale)), "labor_per_unit": num(cell(ws, r, c_labor)) or _ZERO,
            "dia_in": num(cell(ws, r, c_dia)) if c_dia else None, "depth_ft": num(cell(ws, r, c_depth)) if c_depth else None,
            "concrete_cy_total": num(cell(ws, r, c_yds)) if c_yds else None,
            "steel_lb_total": num(cell(ws, r, c_steel)) if c_steel else None,
            "mix": site_mix,
        })
    t.quantity = sum((row["qty"] for row in t.rows), _ZERO)
    t.sale = sum((row["qty"] * (row["unit_sale"] or _ZERO) for row in t.rows), _ZERO) or None


READERS = {
    "mono_slab": read_mono_slab, "slabs": read_slabs, "paving": read_paving, "sidewalk": read_sidewalks,
    "grade_beams": read_beams, "cont_footings": read_beams, "spot_footings": read_footings, "miscellaneous": read_misc,
}


# ------------------------------------------------------------ the book ----


def read_information(wb) -> tuple[dict[str, Any], dict[str, Any]]:
    project: dict[str, Any] = {}
    estimate: dict[str, Any] = {}
    if "Information" not in wb.sheetnames:
        return project, estimate
    ws = wb["Information"]
    labels: dict[str, tuple[int, int]] = {}
    for r in range(1, 25):
        for c in range(1, 14):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                labels.setdefault(norm(v), (r, c))

    def after(label: str, *alts: str):
        for lb in (label, *alts):
            if norm(lb) in labels:
                r, c = labels[norm(lb)]
                return ws.cell(r, c + 1).value
        return None

    project["name"] = (after("JOB NAME:") or "").strip() if isinstance(after("JOB NAME:"), str) else after("JOB NAME:")
    addr = after("JOB ADDRESS:")
    city = after("CITY, STATE:")
    project["location"] = ", ".join(str(x).strip() for x in (addr, city) if x) or None
    bid = after("BID:")
    project["job_number"] = str(bid).strip() if bid else None
    project["gc"] = (after("CLIENT # 1 NAME:", "BID TO:") or None)
    project["contact"] = after("CONTACT:")
    dt = ws.cell(labels[norm("BID:")][0], labels[norm("BID:")][1] + 2).value if norm("BID:") in labels else None
    project["bid_date"] = dt.date() if hasattr(dt, "date") else None
    estimate["name"] = str(after("REVISION:") or "Imported").strip()
    estimate["prepared_by"] = after("PREPARED BY:")
    return project, estimate


def read_pricing(wb) -> dict[str, Decimal]:
    """The Pricing tab's lumber list: name -> unit cost, for the job's price sheet."""
    out: dict[str, Decimal] = {}
    if "Pricing" not in wb.sheetnames:
        return out
    ws = wb["Pricing"]
    hit = None
    for r in range(1, 6):
        for c in range(1, 30):
            if norm(ws.cell(r, c).value) == "UNIT COST":
                hit = (r, c)
    if hit is None:
        return out
    r0, c_cost = hit
    for r in range(r0 + 1, r0 + 80):
        name = ws.cell(r, c_cost - 1).value
        price = num(ws.cell(r, c_cost).value)
        if isinstance(name, str) and name.strip() and price is not None and price > 0:
            out[norm(name)] = price
    return out


def read_workbook(path: str | Path) -> JobSpec:
    """The whole book, as a spec the apply step can act on and a dry run can print."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    project, estimate = read_information(wb)
    spec = JobSpec(path=str(path), project=project, estimate=estimate, lumber=read_pricing(wb))
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        hit = find_label(ws, "Tax Exempt", cols=range(1, 30), rows=range(50, 80))
        if hit is not None:
            spec.project["tax_exempt"] = yes(ws.cell(hit[0], hit[1] + 1).value)
    rates: dict[str, list[Decimal]] = {}
    for name in wb.sheetnames:
        kind = kind_of(name)
        if kind is None:
            continue
        if kind not in READERS_BUILT:
            spec.skipped.append((name, f"a {kind} tab; no reader for it yet"))
            continue
        ws = wb[name]
        t = TabSpec(tab=name, kind=kind, name=name.strip(), unit=UNIT_BY_KIND.get(kind, "SF"))
        READERS[kind](ws, t)
        if not t.rows:
            spec.skipped.append((name, "no priced rows"))
            continue
        money_block(ws, t)
        labor_blocks(ws, t)
        price_row(ws, t)
        spec.tabs.append(t)
        for s in t.supervision:
            key = supervision_key(s["label"])
            if key and s["rate"]:
                rates.setdefault(key, []).append(s["rate"])
    for key, vals in rates.items():
        spec.day_rates[key] = max(set(vals), key=vals.count)
    return spec


# --------------------------------------------------------- the matching ----


def supervision_key(label: str) -> str | None:
    n = norm(label)
    if n.startswith("SUPERINTENDENT"):
        return "labor_super_day_rate"
    if n.startswith("FOREMAN") or n.startswith("FORMAN"):
        return "labor_foreman_day_rate"
    if n.startswith("PROJECT MAN"):
        return "labor_pm_day_rate"
    if n.startswith("EXPENSE"):
        return "labor_expense_day_rate"
    return None


def supervision_code(label: str) -> str | None:
    return {"labor_super_day_rate": "superintendent", "labor_foreman_day_rate": "foreman",
            "labor_pm_day_rate": "pm", "labor_expense_day_rate": "expense"}.get(supervision_key(label) or "")


# A tab labor label -> the rate keys it may be, in order of preference. The
# section's own line set decides which one it reads.
LABOR_KEYS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("FORMING", ("labor_forming_sf", "labor_gb_forming_ff", "labor_forming_ff", "labor_footings_sf")),
    ("GRADING", ("labor_grading_sf",)),
    ("GRADE", ("labor_grading_sf",)),
    ("THICK", ("labor_thick_edge_lf",)),
    ("PLACE", ("labor_place_finish_sf", "labor_place_finish_ff", "labor_place_finish_ea")),
    ("WRECK", ("labor_wreck_sf", "labor_wreck_ff")),
    ("DROPS", ("labor_drops_ff",)),
    ("TIE STEEL", ("labor_tie_steel_ton",)),
    ("CURB", ("labor_curb_lf",)),
    ("RUB", ("labor_rub_patch_sf", "labor_rub_patch_ff")),
    ("EXCAVAT", ("labor_excavate_cy", "labor_excavation_cy")),
    ("BACKFILL", ("labor_backfill_cy",)),
    ("ADA", ("labor_ada_ramp_ea",)),
    ("STAIR", ("labor_stair_tread_lf",)),
    ("HOLD", ("labor_hold_down_ea",)),
    ("BRICK", ("labor_brick_ledge_lf",)),
    ("PILASTER", ("labor_pilasters_ff",)),
)
CONTRACT_KEYS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("PUMP", ("concrete_pump_cy",)),
    ("SAW CUTTING", ("saw_cutting_lf", "joint_soft_cut_lf")),
    ("SAW JOINT", ("joint_control_lf",)),
    ("EXPANSION JOINT", ("joint_construction_lf",)),
    ("CONTROL JOINT", ("joint_control_lf",)),
    ("HAUL OFF", ("haul_off_cy",)),
    ("CURB DEMO", ("demo_lf",)),
    ("DEMO", ("demo_lf",)),
    ("CURE", ("cure_sf",)),
    ("STAMP", ("stamping_sf",)),
    ("COLOR", ("integral_color_cy",)),
    ("ACID", ("acid_etch_sf",)),
    ("SANDBLAST", ("acid_etch_sf",)),
    ("WATER PROOF", ("waterproofing_sf",)),
    ("WATERPROOF", ("waterproofing_sf",)),
    ("SLIP", ("slip_form_sf",)),
)
EQUIPMENT_CODES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("SKY TRACK", ("skytrack",)), ("SKYTRACK", ("skytrack",)), ("MINI EX", ("mini_excavator", "backhoe")),
    ("TRENCHER", ("trencher",)), ("SKID", ("skid_steer", "bobcat")), ("BOB CAT", ("bobcat", "skid_steer")),
    ("BOBCAT", ("bobcat", "skid_steer")), ("BACK HOE", ("backhoe", "mini_excavator")), ("BACKHOE", ("backhoe", "mini_excavator")),
    ("COMPACTOR", ("compactor",)), ("VAULT", ("vault",)), ("STORAGE", ("storage",)), ("FORK", ("fork_truck",)),
    ("LIGHT TOWER", ("light_tower",)), ("CRANE", ("crane",)), ("MISC", ("misc_equip",)),
)
CONTRACT_CODES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("PUMP", ("concrete_pump",)), ("SAW CUTTING", ("saw_cutting", "soft_cut")), ("SAW JOINT", ("saw_joint_sealant",)),
    ("EXPANSION JOINT", ("joint_construction",)), ("CONTROL JOINT", ("joint_control",)), ("HAUL OFF", ("haul_off",)),
    ("CURB DEMO", ("demo",)), ("CURE", ("cure",)), ("STAMP", ("stamping",)), ("COLOR", ("integral_color",)),
    ("ACID", ("acid_etch",)), ("SANDBLAST", ("acid_etch",)), ("WATER PROOF", ("waterproofing",)), ("WATERPROOF", ("waterproofing",)),
    ("SLIP", ("slip_forming",)), ("OUT OF TOWN", ("out_of_town",)), ("MISC", ("misc_contract",)),
)
# A tab labor label -> the app's line codes it may be, so a line the tab does not carry can be told apart.
LABOR_CODES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("FORMING", ("forming", "gb_forming", "footings")), ("GRADING", ("grading", "layout")), ("GRADE", ("grading", "layout")),
    ("PLACE", ("place_finish",)), ("WRECK", ("wreck", "cleanup")), ("DROPS", ("drops",)), ("TIE STEEL", ("tie_steel", "rebar")),
    ("CURB", ("curb",)), ("RUB", ("rub_patch",)), ("EXCAVAT", ("excavate", "excavation")), ("BACKFILL", ("backfill",)),
    ("ADA", ("ada_ramps",)), ("STAIR", ("stair_treads",)), ("THICK", ("thick_edge",)), ("PILASTER", ("pilasters",)),
    ("HOLD", ("hold_downs",)), ("CABLE", ("cable_placement",)), ("EXTRA", ("extra_hours",)), ("LABOR ADD", ("labor_add",)),
    ("LABOR ADJ", ("labor_add",)), ("BRICK", ("brick_ledge",)), ("EDGE", ("edge_rails",)), ("FOOTING", ("footings",)),
    ("FRENCH", ("french_drains",)), ("PIER CAP", ("pier_caps",)), ("RESHOR", ("reshoring",)), ("STUD", ("stud_rails",)),
)
# Labor lines the app cannot derive a quantity for: the tab's count goes on the line.
TYPED_LABOR_CODES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("ADA", ("ada_ramps",)), ("STAIR", ("stair_treads",)), ("HOLD", ("hold_downs",)), ("EXTRA", ("extra_hours",)),
    ("LABOR ADD", ("labor_add",)), ("LABOR ADJ", ("labor_add",)),
)
_AREA = ("SF", "SQ FT", "SQ. FT", "SQFT", "/SQ FT", "SQ FT.", "FF", "FC FT", "FC.FT", "FACE", "/FC FT", "SQ FF", "/FC. FT")
UNIT_WORDS = {
    "SF": _AREA, "FF": _AREA,
    "LF": ("LF", "LN FT", "/LN FT", "LN. FT"), "TON": ("TON", "/TON", "TONS"),
    "CY": ("CY", "CU YD", "CU. YD", "/YD", "/ YD", "/YRD", "YRD", "/YARD", "YARD", "YDS"),
    "EA": ("EA", "EACH", "/EACH", "QTY", "QNTY"), "DAY": ("DAY", "DAYS", "/DAY"), "LS": ("LS", "FOR", "EA"),
}


def unit_matches(key_unit: str | None, tab_unit: Any) -> bool:
    if not key_unit or tab_unit in (None, ""):
        return True
    u = norm(tab_unit).replace(".", "")
    return any(u == norm(w).replace(".", "") or u.startswith(norm(w).replace(".", "")) for w in UNIT_WORDS.get(key_unit, (key_unit,)))


def first_code(label: str, table, present=None) -> str | None:
    """The app's line code a tab label names — the first candidate the section carries, when `present` is given."""
    n = norm(label)
    for needle, codes in table:
        if n.startswith(needle) or needle in n:
            for code in codes:
                if present is None or code in present:
                    return code
            return None
    return None


def rate_key_for(label: str, table, read_keys: set[str], tab_unit: Any = None) -> str | None:
    """The first key the section reads that the label names and whose unit the tab's agrees with."""
    from app.services.price_book import MONETARY_KEYS

    n = norm(label)
    for needle, keys in table:
        if n.startswith(needle) or needle in n:
            for k in keys:
                if k in read_keys and unit_matches(MONETARY_KEYS.get(k, ("", None))[1], tab_unit):
                    return k
            return None
    return None


# ---------------------------------------------------------- the apply ----


MIX_CODE_BY_STRENGTH = {3000: "3000-SC", 3500: "3500-SC", 4000: "4000-SC", 4500: "4500-SC", 5000: "5000-SC"}


def app_mix_code(name: str) -> str | None:
    """The catalog code the tab's mix name means: strength, and ash where the name says so."""
    m = re.search(r"(\d{4})", name or "")
    if not m:
        return None
    psi = int(m.group(1))
    code = MIX_CODE_BY_STRENGTH.get(psi)
    if code and "ASH" in norm(name):
        code = code.replace("-SC", "-ASH")
    return code


REBAR_MATERIAL_BY_KIND = {
    "mono_slab": "REBAR PIERS / PT slabs", "slabs": "REBAR PIERS / PT slabs",
    "grade_beams": "REBAR GRADE BEAM", "cont_footings": "REBAR GRADE BEAM", "spot_footings": "REBAR GRADE BEAM",
    "paving": "REBAR PAVING", "sidewalk": "REBAR PAVING",
}
MISC_LIBRARY_ALIASES: tuple[tuple[str, str], ...] = (
    ("LIGHT POLE", "Light Pole Bases"), ("PIPE BOLLARD", "Pipe Bollards (Install Only)"),
    ("BOLLARD", "Bollard / Landscape Light Bases"), ("ELEVATOR PIT", "Elevator Pits"),
    ("TRANSFORME", "Transformer Pads (Three Phase)"), ("GATE TRACK", "Gate Track"), ("ADA RAMP", "ADA Ramps"),
    ("STAIR TREAD", "Stair Treads"), ("BIKE RACK", "Bike Rack (Installation)"), ("TRELLIS", "Trellis Ftgs"),
    ("FIRE PIT", "Fire Pit Footings"), ("GRILL", "Grill Footing (12\" thick 369 sq ft)"),
    ("MONUMENT", "Monument Base (15' x 3' x 36\")"), ("RADON", "Radon Pits"),
)


@dataclass
class Report:
    project_id: Any = None
    estimate_id: Any = None
    sections: list[dict[str, Any]] = field(default_factory=list)
    prices: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def tab_sale(self) -> Decimal:
        return sum((s["tab_sale"] or _ZERO for s in self.sections), _ZERO)

    @property
    def app_sale(self) -> Decimal:
        return sum((s["app_sale"] or _ZERO for s in self.sections), _ZERO)


def _mix_ids(db: Session) -> dict[str, int]:
    return {code: mid for code, mid in db.execute(text("SELECT code, id FROM mix_designs WHERE is_active")).all()}


def _material_id(db: Session, name: str) -> int | None:
    return db.execute(
        text("SELECT id FROM materials WHERE regexp_replace(upper(name), '\\s+', ' ', 'g') = :n LIMIT 1"),
        {"n": norm(name)},
    ).scalar()


def _set_sheet(db: Session, estimate_id, kind: str, *, ref_id: int | None = None, ref_key: str | None = None,
               value: Decimal, note: str, report: Report) -> bool:
    """What this job pays for one item. A row the pull never made — a mix the catalog has no price for — is added."""
    from datetime import datetime, timezone

    from app.models.estimate_price import EstimatePrice
    from app.services.price_book import EQUIPMENT_CATEGORY, MONETARY_KEYS, RATE_CATEGORY, set_price, sheet_rows

    for p in sheet_rows(db, estimate_id):
        if p.kind == kind and (ref_id is None or p.ref_id == ref_id) and (ref_key is None or p.ref_key == ref_key):
            if p.value is not None and Decimal(str(p.value)) == value:
                return True
            set_price(db, p, value=value, note=note)
            report.prices.append(f"{p.label or ref_key or ref_id}: {value}")
            return True
    label = unit = category = None
    if kind == "mix":
        row = db.execute(text("SELECT code, unit FROM mix_designs WHERE id = :i"), {"i": ref_id}).first()
        label, unit, category = (row[0], row[1], "concrete") if row else (str(ref_id), "CY", "concrete")
    elif kind == "material":
        row = db.execute(text("SELECT name, unit, category FROM materials WHERE id = :i"), {"i": ref_id}).first()
        label, unit, category = (row[0], row[1], row[2]) if row else (str(ref_id), None, None)
    elif kind == "equipment":
        row = db.execute(text("SELECT name, unit FROM equipment WHERE id = :i"), {"i": ref_id}).first()
        label, unit, category = (row[0], row[1] or "DAY", EQUIPMENT_CATEGORY) if row else (str(ref_id), "DAY", EQUIPMENT_CATEGORY)
    elif kind == "setting":
        label, unit = MONETARY_KEYS.get(ref_key or "", (ref_key, None))
        category = RATE_CATEGORY
    if label is None:
        report.notes.append(f"price sheet has no {kind} row for {ref_key or ref_id}")
        return False
    db.add(EstimatePrice(
        estimate_id=estimate_id, kind=kind, scope=None, ref_id=ref_id, ref_key=ref_key, label=label, unit=unit,
        category=category, catalog_value=None, value=value.quantize(_Q4), is_edited=True, note=note,
        pulled_at=datetime.now(timezone.utc),
    ))
    db.flush()
    report.prices.append(f"{label}: {value} (no master price; the workbook's)")
    return True


def _section_rate(db: Session, section_id, key: str, value: Decimal, note: str) -> None:
    db.execute(
        text(
            "INSERT INTO section_rates (section_id, key, value, note) VALUES (:s, :k, :v, :n) "
            "ON CONFLICT (section_id, key) DO UPDATE SET value = excluded.value, note = excluded.note, updated_at = now()"
        ),
        {"s": str(section_id), "k": key, "v": value, "n": note},
    )


def _switch_forming(db: Session, section, code: str) -> bool:
    """Switch a forming line off the way the screen does; False when the set has no such line."""
    from app.models.estimate_forming import EstimateFormingLine
    from app.services.forming import set_forming_line_enabled

    row = db.scalars(select(EstimateFormingLine).where(
        EstimateFormingLine.section_id == section.id, EstimateFormingLine.code == code)).first()
    if row is None or not row.enabled:
        return False
    set_forming_line_enabled(db, section.id, code, False)
    return True


def _pin_days(db: Session, section, code: str, days: Decimal, rate: Decimal | None, *, equipment: bool) -> bool:
    """Type the days (and, on a machine, the rate) onto a line the way the screen does."""
    if equipment:
        from app.services.estimate_equipment import update_equipment_line

        try:
            update_equipment_line(db, section.id, code, days_qty=days, rate=rate, mark_manual=True)
        except ValueError:
            return False
        return True
    from app.services.labor import update_labor_line

    try:
        update_labor_line(db, section.id, code, qty=days, mark_manual=True)
    except ValueError:
        return False
    return True


def apply(db: Session, spec: JobSpec, *, replace: bool = False, project_name: str | None = None,
          estimate_name: str | None = None) -> Report:
    """Write the spec onto the database through the app's services. Commits."""
    from app.models.estimate import Estimate
    from app.models.estimate_section import EstimateSection
    from app.models.project import Project
    from app.services import section_rates as sr
    from app.services.costing import refresh_estimate_totals, refresh_pour_costs
    from app.services.estimate_equipment import get_or_refresh_equipment
    from app.services.forming import get_or_refresh_forming
    from app.services.labor import get_or_refresh_labor
    from app.services.price_book import pull_prices
    from app.services.recalc import recalc_section

    report = Report()
    pname = (project_name or spec.project.get("name") or "Imported job").strip()
    ename = (estimate_name or spec.estimate.get("name") or "Imported").strip()

    project = db.scalars(select(Project).where(Project.name == pname)).first()
    if project is None:
        project = Project(
            name=pname, job_number=spec.project.get("job_number"), gc=spec.project.get("gc"),
            location=spec.project.get("location"), bid_date=spec.project.get("bid_date"),
            tax_exempt=bool(spec.project.get("tax_exempt", False)),
        )
        db.add(project)
        db.flush()
        report.notes.append(f"project created: {pname}")
    else:
        report.notes.append(f"project found: {pname}")
    report.project_id = project.id

    existing = db.scalars(select(Estimate).where(Estimate.project_id == project.id, Estimate.name == ename)).first()
    if existing is not None:
        if not replace:
            raise RuntimeError(f"estimate {ename!r} already exists on {pname!r}; pass --replace to rebuild it")
        db.delete(existing)
        db.flush()
        report.notes.append(f"estimate replaced: {ename}")
    margins = [t.margin for t in spec.tabs if t.margin is not None]
    estimate = Estimate(
        project_id=project.id, name=ename, status="draft", version=1,
        margin_pct=max(set(margins), key=margins.count) if margins else Decimal("0.20"),
        contingency_pct=Decimal("0"),
        notes=f"Imported from {Path(spec.path).name} on {date.today().isoformat()}",
    )
    db.add(estimate)
    db.flush()
    pull_prices(db, estimate.id)
    report.estimate_id = estimate.id

    # The job's price sheet: the mixes, the steel, the PT, the sand, the lumber, the day rates.
    mix_ids = _mix_ids(db)
    stamp = f"{Path(spec.path).name}"
    mix_prices: dict[str, Decimal] = {}
    for t in spec.tabs:
        for name, price in t.mixes.values():
            code = app_mix_code(name)
            if code and price and code in mix_ids:
                mix_prices.setdefault(code, price)
    for code, price in mix_prices.items():
        _set_sheet(db, estimate.id, "mix", ref_id=mix_ids[code], value=price, note=stamp, report=report)
    steel: dict[str, list[Decimal]] = {}
    pt: list[Decimal] = []
    sand: list[Decimal] = []
    for t in spec.tabs:
        mat = REBAR_MATERIAL_BY_KIND.get(t.kind)
        if mat and t.prices.get("steel_lb"):
            steel.setdefault(mat, []).append(t.prices["steel_lb"])
        if t.prices.get("pt_sf") and t.kind == "mono_slab":
            pt.append(t.prices["pt_sf"])
        if t.prices.get("sand_cy"):
            sand.append(t.prices["sand_cy"])
    for mat, vals in steel.items():
        mid = _material_id(db, mat)
        if mid:
            _set_sheet(db, estimate.id, "material", ref_id=mid, value=max(set(vals), key=vals.count), note=stamp, report=report)
    if pt and (mid := _material_id(db, "POST TENSION CABLES")):
        _set_sheet(db, estimate.id, "material", ref_id=mid, value=max(set(pt), key=pt.count), note=stamp, report=report)
    if sand and (mid := _material_id(db, "SAND DELIVERED PER CY")):
        _set_sheet(db, estimate.id, "material", ref_id=mid, value=max(set(sand), key=sand.count), note=stamp, report=report)
    for name, price in spec.lumber.items():
        mid = _material_id(db, name)
        if mid:
            _set_sheet(db, estimate.id, "material", ref_id=mid, value=price, note=f"Pricing tab, {stamp}", report=report)
    for key, rate in spec.day_rates.items():
        _set_sheet(db, estimate.id, "setting", ref_key=key, value=rate, note=stamp, report=report)
    # The pump is priced off the equipment catalog on the slab tabs; the tabs agree on a rate.
    pumps = [c["rate"] for t in spec.tabs for c in t.contract
             if "PUMP" in norm(c["label"]) and c.get("rate") and c.get("cost") and c.get("enabled", True)]
    if pumps:
        pid = db.execute(text("SELECT id FROM equipment WHERE upper(name) LIKE 'CONCRETE PUMP%' AND is_active ORDER BY id LIMIT 1")).scalar()
        if pid:
            _set_sheet(db, estimate.id, "equipment", ref_id=pid, value=max(set(pumps), key=pumps.count), note=stamp, report=report)
    db.flush()

    for order, t in enumerate(spec.tabs):
        section = EstimateSection(
            estimate_id=estimate.id, kind=t.kind, name=t.name, unit=t.unit, sort_order=(order + 1) * 10,
            margin_pct=t.margin if t.margin is not None else estimate.margin_pct, contingency_pct=Decimal("0"),
            labor_subcontracted=(sum(1 for ln in t.labor if ln["sub"]) * 2 >= len(t.labor)) if t.labor else True,
            waste_concrete=t.waste.get("concrete"), waste_rebar=t.waste.get("rebar"), waste_sand=t.waste.get("sand"),
            tax_exempt=None,
            notes=f"Imported from the {t.tab!r} tab of {Path(spec.path).name}",
        )
        db.add(section)
        db.flush()
        sr.seed(db, section)
        entry = {"tab": t.tab, "kind": t.kind, "section_id": section.id, "rows": 0, "tab_sale": t.sale,
                 "tab_cost": t.total_cost, "tab_quantity": t.quantity, "typed": [], "rates": [], "off": [], "unmatched": []}
        _rows_for(db, section, t, mix_ids, entry)
        # A slab tab with no poly line prices no vapor barrier; the app would add one.
        if t.kind in ("mono_slab", "slabs") and not any("POLY" in norm(n) or "VISQUEEN" in norm(n) or "STEGO" in norm(n)
                                                        for n in (m[0] for m in t.mixes.values())):
            _section_rate(db, section.id, "vapor_barrier_enabled", _ZERO, f"the {t.tab!r} tab prices no vapor barrier")
            entry["rates"].append("vapor_barrier_enabled=0")
        db.flush()
        recalc_section(db, section)
        db.flush()
        get_or_refresh_forming(db, section.id)
        get_or_refresh_labor(db, section.id)
        get_or_refresh_equipment(db, section.id)
        db.flush()
        for code, on in t.switches.items():
            if not on and _switch_forming(db, section, code):
                entry["off"].append(f"{code} (the tab's N)")
        _rates_for(db, section, t, entry)
        db.flush()
        recalc_section(db, section)
        db.flush()
        _days_for(db, section, t, entry)
        db.flush()
        refresh_pour_costs(db, section)
        db.flush()
        db.refresh(section)
        entry["app_sale"] = Decimal(str(section.calc_total_sale or 0))
        entry["app_cost"] = Decimal(str(section.calc_total_cost or 0))
        entry["app_quantity"] = section.calc_quantity
        report.sections.append(entry)
    refresh_estimate_totals(db, estimate)
    db.commit()
    return report


def _rows_for(db: Session, section, t: TabSpec, mix_ids: dict[str, int], entry: dict[str, Any]) -> None:
    from app.models.beam_run import BeamRun
    from app.models.beam_type import EstimateBeamType
    from app.models.grade_beam import GradeBeam
    from app.models.misc_item import MiscItem
    from app.models.mono_slab import MonoSlab
    from app.models.wall_run import WallRun
    from app.services.beams import refresh_beam_run_calcs
    from app.services.calc import refresh_mono_slab_calcs
    from app.services.misc import refresh_misc_item_calcs, seed_from_library
    from app.services.walls import refresh_wall_run_calcs

    def mix_id(n: int | None) -> int | None:
        if n is None or n not in t.mixes:
            return None
        code = app_mix_code(t.mixes[n][0])
        return mix_ids.get(code) if code else None

    if t.kind in ("mono_slab", "slabs", "paving", "sidewalk"):
        types: dict[tuple[int, str], EstimateBeamType] = {}
        if t.kind == "mono_slab":
            roles: dict[int, set[str]] = {}
            for uses in t.usages.values():
                for kind, n, _ in uses:
                    roles.setdefault(n, set()).add(kind)
            for i, bt in enumerate(t.beam_types):
                for role in sorted(roles.get(bt["n"], {"grade_beam"})):
                    row = EstimateBeamType(
                        section_id=section.id, label=f"GB {bt['n']}" + ("" if role == "grade_beam" else f" ({role})"),
                        kind=role, width_in=bt["width_in"], height_in=bt["height_in"],
                        top_bars_count=bt["top_bars_count"], top_bars_size=bt["top_bars_size"],
                        bottom_bars_count=bt["bottom_bars_count"], bottom_bars_size=bt["bottom_bars_size"],
                        mid_bars_count=bt["mid_bars_count"], mid_bars_size=bt["mid_bars_size"],
                        stirrup_size=bt["stirrup_size"], stirrup_spacing_in=bt["stirrup_spacing_in"],
                        l_bars_size=bt["l_bars_size"], l_bars_spacing_in=bt["l_bars_spacing_in"],
                        l_bars_length_ft=bt["l_bars_length_ft"], sort_order=(i + 1) * 10,
                    )
                    db.add(row)
                    types[(bt["n"], role)] = row
            db.flush()
        for i, r in enumerate(t.rows):
            fields = {k: v for k, v in r.items() if k not in ("mix", "carton")}
            slab = MonoSlab(section_id=section.id, mix_design_id=mix_id(r.get("mix")), sort_order=(i + 1) * 10, **fields)
            db.add(slab)
            db.flush()
            for j, (kind, n, lf) in enumerate(t.usages.get(i, [])):
                bt = types.get((n, kind))
                if bt is None:
                    entry["unmatched"].append(f"{r['description']}: beam type {n} ({kind}) not in the schedule")
                    continue
                # A usage's kind is its type's (the property has no setter); the type was made per role above.
                db.add(GradeBeam(mono_slab_id=slab.id, beam_type_id=bt.id, length_lf=lf, sort_order=(j + 1) * 10))
            db.flush()
            refresh_mono_slab_calcs(db, slab, section)
            entry["rows"] += 1
    elif t.kind in ("grade_beams", "cont_footings"):
        for i, r in enumerate(t.rows):
            fields = {k: v for k, v in r.items() if k != "mix"}
            run = BeamRun(section_id=section.id, mix_design_id=mix_id(r.get("mix")), sort_order=(i + 1) * 10, **fields)
            db.add(run)
            db.flush()
            refresh_beam_run_calcs(db, run, section)
            entry["rows"] += 1
    elif t.kind == "spot_footings":
        for i, r in enumerate(t.rows):
            fields = {k: v for k, v in r.items() if k not in ("mix", "pilaster_count")}
            run = WallRun(
                section_id=section.id, footing_mix_design_id=mix_id(r.get("mix")), mix_design_id=mix_id(r.get("mix")),
                sort_order=(i + 1) * 10, backfill=False, wall_thick_in=_ZERO, wall_height_in=_ZERO,
                length_ft=(r["footing_each_ft"] or _ZERO) * r["footing_count"], weld_plate=False, **fields,
            )
            db.add(run)
            db.flush()
            refresh_wall_run_calcs(db, run, section)
            entry["rows"] += 1
    elif t.kind == "miscellaneous":
        seed_from_library(db, section)
        items = {norm(it.description): it for it in db.scalars(select(MiscItem).where(MiscItem.section_id == section.id)).all()}
        for i, r in enumerate(t.rows):
            target = None
            n = norm(r["description"])
            for needle, lib_name in MISC_LIBRARY_ALIASES:
                if needle in n and norm(lib_name) in items:
                    target = items[norm(lib_name)]
                    break
            if target is None and n in items:
                target = items[n]
            if target is None:
                # A pad on the tab is an area (SF) and a thickness (in) in the DIA and DEPTH
                # cells — the app's slab shape; anything else is a round base, dia and depth.
                pad = "PAD" in n
                dia, depth = r.get("dia_in"), r.get("depth_ft")
                target = MiscItem(
                    section_id=section.id, code=None, description=r["description"], unit="EA",
                    shape="slab" if pad else "round", qty=_ZERO, pours_concrete=True,
                    dim_a=dia, dim_b=depth, dim_c=None, sort_order=900 + i,
                    steel_lb_per_unit=((r["steel_lb_total"] / r["qty"]).quantize(Decimal("0.001"))
                                       if r.get("steel_lb_total") and r["qty"] else _ZERO),
                )
                db.add(target)
                db.flush()
                entry["unmatched"].append(f"{r['description']}: not in the library; added as a {target.shape}")
            else:
                if r.get("dia_in") and target.shape == "round":
                    target.dim_a = r["dia_in"]
                if r.get("depth_ft") and target.shape == "round":
                    target.dim_b = r["depth_ft"]
            target.qty = r["qty"]
            target.unit_sale = r["unit_sale"]
            target.labor_per_unit = r["labor_per_unit"]
            target.subcontracted = r["subcontracted"]
            if mix_id(r.get("mix")):
                target.mix_design_id = mix_id(r.get("mix"))
            db.flush()
            refresh_misc_item_calcs(db, target, section)
            entry["rows"] += 1
    db.flush()


def _rates_for(db: Session, section, t: TabSpec, entry: dict[str, Any]) -> None:
    """The tab's labor and contract rates onto the section, where the section reads such a key."""
    from app.services import section_rates as sr
    from app.services.price_book import MONETARY_KEYS

    _ = MONETARY_KEYS
    read = set(sr.keys_read(db, section))
    note = f"the {t.tab!r} tab"
    for ln in t.labor:
        if ln["rate"] is None:
            continue
        key = rate_key_for(ln["label"], LABOR_KEYS, read, ln.get("unit"))
        if key is None:
            if ln.get("cost"):
                entry["unmatched"].append(f"labor {ln['label']!r} @ {ln['rate']} {ln.get('unit') or ''}: no such rate on this section")
            continue
        _section_rate(db, section.id, key, ln["rate"], note)
        entry["rates"].append(f"{key}={ln['rate']}")
    from app.models.estimate_equipment import EstimateEquipmentLine

    lines = {r.code: r for r in db.scalars(select(EstimateEquipmentLine).where(EstimateEquipmentLine.section_id == section.id)).all()}
    for ln in t.contract:
        if ln.get("rate") is None or not ln.get("cost") or not ln.get("enabled", True):
            continue   # a service the tab charges nothing for is switched off below, not rated
        key = rate_key_for(ln["label"], CONTRACT_KEYS, read, ln.get("unit"))
        if key is None:
            code = first_code(ln["label"], CONTRACT_CODES, lines)
            if code is not None and Decimal(str(lines[code].rate or 0)) == ln["rate"]:
                continue   # priced off the job's equipment sheet at the tab's own number
            entry["unmatched"].append(f"contract {ln['label']!r} @ {ln['rate']} {ln.get('unit') or ''}: no such rate on this section"
                                      + (f"; the app's line is at {lines[code].rate}" if code is not None else ""))
            continue
        _section_rate(db, section.id, key, ln["rate"], note)
        entry["rates"].append(f"{key}={ln['rate']}")


def _days_for(db: Session, section, t: TabSpec, entry: dict[str, Any]) -> None:
    """The tab's supervision days and its machines' days, typed onto the lines; a pump turned off stays off."""
    from app.models.estimate_equipment import EstimateEquipmentLine
    from app.services.estimate_equipment import update_equipment_line

    for s in t.supervision:
        code = supervision_code(s["label"])
        if code is None or s["days"] is None:
            continue
        if _pin_days(db, section, code, s["days"], s["rate"], equipment=False):
            entry["typed"].append(f"{code} {s['days']} days")
        else:
            entry["unmatched"].append(f"supervision {s['label']!r}: no such line on this section")
    # A labor line the app cannot derive — ADA ramps, stair treads, hold-downs — takes the tab's count.
    from app.models.estimate_labor import EstimateLaborLine

    from app.services.labor import update_labor_line

    labor_lines = {r.code: r for r in db.scalars(select(EstimateLaborLine).where(EstimateLaborLine.section_id == section.id)).all()}
    kept: set[str] = set()
    for ln in t.labor:
        code = first_code(ln["label"], LABOR_CODES, labor_lines)
        if code is None:
            continue
        if ln.get("cost"):
            kept.add(code)
        typed = first_code(ln["label"], TYPED_LABOR_CODES, labor_lines)
        if typed and ln.get("qty") and ln.get("cost") and Decimal(str(labor_lines[typed].qty or 0)) == 0:
            if _pin_days(db, section, typed, ln["qty"], None, equipment=False):
                entry["typed"].append(f"{typed} {ln['qty']} {ln.get('unit') or ''}".rstrip())
    # A pour's add $/SF lands on LABOR ADD whether or not the tab's block lists the line.
    if any(r.get("paving_add_per_sf") for r in t.rows):
        kept.add("labor_add")
    # A field labor line the tab does not carry (or carries at nothing) is switched off, and said so.
    for code, row in labor_lines.items():
        if row.group_name == "labor" and row.enabled and code not in kept and Decimal(str(row.ext_cost or 0)) != 0:
            try:
                update_labor_line(db, section.id, code, enabled=False, mark_manual=None)
                entry["off"].append(f"labor {code} ({row.ext_cost})")
            except ValueError:
                pass

    lines = {r.code: r for r in db.scalars(select(EstimateEquipmentLine).where(EstimateEquipmentLine.section_id == section.id)).all()}
    kept_equipment: set[str] = set()

    def switch_off(code: str, why: str) -> None:
        try:
            update_equipment_line(db, section.id, code, enabled=False, mark_manual=False)
            entry["off"].append(f"{code} ({why})")
        except ValueError:
            pass

    for e in t.equipment:
        code = first_code(e["label"], EQUIPMENT_CODES, lines)
        if code is None:
            if e.get("cost"):
                entry["unmatched"].append(f"equipment {e['label']!r} {e['days']} days @ {e['rate']} = {e['cost']:.2f}: no such line on this section")
            continue
        if not e.get("cost") or not e.get("days"):
            continue
        kept_equipment.add(code)
        if _pin_days(db, section, code, e["days"], e["rate"], equipment=True):
            entry["typed"].append(f"{code} {e['days']} days" + (f" @ {e['rate']}" if e["rate"] else ""))
    said_no: set[str] = set()
    for c in t.contract:
        code = first_code(c["label"], CONTRACT_CODES, lines)
        if code is None or code in kept_equipment:
            continue
        if not c.get("enabled", True):
            said_no.add(code)   # "Pump Paving?  N": off whatever it would cost
        elif c.get("cost"):
            kept_equipment.add(code)
            if code in ("out_of_town", "misc_contract") and c.get("qty"):
                if _pin_days(db, section, code, c["qty"], c.get("rate"), equipment=True):
                    entry["typed"].append(f"{code} {c['qty']} @ {c.get('rate')}")
    for code, row in lines.items():
        if row.enabled and code not in kept_equipment and (Decimal(str(row.ext_cost or 0)) != 0 or code in said_no):
            switch_off(code, "the tab's N" if code in said_no else f"{row.ext_cost}, not on the tab")


# ---------------------------------------------------------- the print ----


def describe(spec: JobSpec) -> str:
    out = [f"{spec.path}", f"project: {spec.project}", f"estimate: {spec.estimate}",
           f"day rates: {spec.day_rates}", f"lumber prices: {len(spec.lumber)}"]
    for t in spec.tabs:
        out.append(f"\n== {t.tab} -> {t.kind} ({t.unit}); {len(t.rows)} rows; qty {t.quantity}; sale {t.sale}; "
                   f"margin {t.margin}; tax exempt {t.tax_exempt}; supplier {t.supplier}")
        out.append(f"   mixes {t.mixes}")
        out.append(f"   prices {t.prices} waste {t.waste}")
        if t.beam_types:
            out.append(f"   beam types: {[(b['n'], str(b['width_in']), str(b['height_in'])) for b in t.beam_types]}")
        for i, r in enumerate(t.rows[:12]):
            out.append(f"   row {i}: " + ", ".join(f"{k}={v}" for k, v in r.items() if v not in (None, False, 0)))
            if i in t.usages:
                out.append(f"          beams: {t.usages[i]}")
        if len(t.rows) > 12:
            out.append(f"   ... {len(t.rows) - 12} more rows")
        out.append("   labor: " + "; ".join(f"{ln['label']} {ln['rate']} {ln.get('unit')}" for ln in t.labor if ln["rate"]))
        out.append("   supervision: " + "; ".join(f"{s['label']} {s['days']}d @ {s['rate']}" for s in t.supervision if s["days"]))
        out.append("   equipment: " + "; ".join(f"{e['label']} {e['days']}d @ {e['rate']}" for e in t.equipment if e["days"]))
        out.append("   contract: " + "; ".join(f"{c['label']} rate {c.get('rate')} qty {c.get('qty')}{'' if c.get('enabled', True) else ' OFF'}" for c in t.contract if c.get("rate") or c.get("qty")))
        for n in t.notes:
            out.append(f"   note: {n}")
    for name, why in spec.skipped:
        out.append(f"skipped {name!r}: {why}")
    return "\n".join(out)


def tie_out(report: Report) -> str:
    out = [f"{'section':40} {'tab sale':>14} {'app sale':>14} {'diff':>12} {'%':>7}   {'tab cost':>13} {'app cost':>13}"]
    for s in report.sections:
        ts, aps = s["tab_sale"] or _ZERO, s["app_sale"] or _ZERO
        diff = aps - ts
        pct = (diff / ts * 100) if ts else _ZERO
        out.append(f"{s['tab'][:40]:40} {ts:>14,.2f} {aps:>14,.2f} {diff:>12,.2f} {pct:>6.1f}%   {(s['tab_cost'] or _ZERO):>13,.2f} {s['app_cost']:>13,.2f}")
    diff = report.app_sale - report.tab_sale
    pct = (diff / report.tab_sale * 100) if report.tab_sale else _ZERO
    out.append(f"{'TOTAL':40} {report.tab_sale:>14,.2f} {report.app_sale:>14,.2f} {diff:>12,.2f} {pct:>6.1f}%")
    for s in report.sections:
        out.append(f"\n{s['tab']}: {s['rows']} rows; quantity tab {s['tab_quantity']} / app {s['app_quantity']}")
        if s["rates"]:
            out.append("  rates: " + ", ".join(s["rates"]))
        if s["typed"]:
            out.append("  typed: " + ", ".join(s["typed"]))
        if s["off"]:
            out.append("  off:   " + ", ".join(s["off"]))
        for u in s["unmatched"]:
            out.append(f"  ! {u}")
    for n in report.notes:
        out.append(f"  - {n}")
    if report.prices:
        out.append("price sheet: " + "; ".join(report.prices))
    return "\n".join(out)
