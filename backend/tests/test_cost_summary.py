"""
The estimate Summary (sql/088): the workbook's Summary tab, read off the job.

Chad, 2026-09-09: "there is something missing that the office uses from the
estimates.. a summary like that is in the excel spreadsheet.. pulls all the
materials form each section, supervision, everything thats in a project.. a
total of each the a breakdown per section.." — and "build it".

What is pinned here:

  * every line every shape writes has a cost code — the matrix, one shape
    per test, so a new line with no home is a red test and not an
    UNASSIGNED column on a summary somebody hands the office;
  * a section's columns ADD UP TO ITS STORED COST, to the cent, with the
    per-row rounding named as `difference` rather than hidden; the codes
    plus the tax do the same; profit plus contingency plus cost is the sale;
  * the columns mean what the tab's mean: concrete is the concrete purchase,
    supervision is the superintendent, foreman and expense lines, PM is the
    PM line, the pump is an OTHER SUB, out of town is its own column, the
    lumber block files under Lumber and the chairs under accessories;
  * the section's Y/N moves labor between LABOR and SUB LABOR and between
    the in-house and sub-labor codes, without moving a dollar;
  * the job adds its sections; the lower block is the tab's L57:L63;
  * a senior refiles a line and every summary follows; an estimator may not;
  * a line the table does not know is counted as UNASSIGNED and named;
  * a section nobody has opened says so;
  * the .xlsx is the tab.
"""

from __future__ import annotations

import importlib
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.estimate_equipment import EstimateEquipmentLine
from app.models.estimate_forming import EstimateFormingLine
from app.models.estimate_labor import EstimateLaborLine
from app.services.cost_summary import (
    COLUMN_KEYS,
    LABOR_INSURANCE_PCT,
    MATERIAL_COLUMNS,
    OTHER_COLUMNS,
    SUB_LABOR_COLUMNS,
    load_chart,
    section_summary,
)
from app.services.costing import cost_units, refresh_estimate_totals, refresh_pour_costs
from app.services.material_costs import section_material_costs
from tests import mono_slab_fixture as mf

D = Decimal
SUBTOTALS = {"000040", "000060"}
SHAPES = [
    "mono_slab_fixture", "paving_fixture", "sidewalks_fixture", "piers_fixture", "grade_beams_fixture",
    "walls_fixture", "spot_footings_fixture", "columns_fixture", "slabs_fixture", "deck_fixture",
    "slab_on_deck_fixture", "panels_fixture", "misc_fixture",
]


def _full(db, estimate, mod_name):
    """A section with its three line sets built and its costs allocated — test_write_paths' recipe."""
    from app.services.estimate_equipment import get_or_refresh_equipment
    from app.services.forming import get_or_refresh_forming
    from app.services.labor import get_or_refresh_labor
    from app.services.recalc import recalc_section

    mod = importlib.import_module(f"tests.{mod_name}")
    section = mod.build(db, estimate)
    if hasattr(mod, "type_the_supervision"):
        mod.type_the_supervision(db, section.id)
    db.flush()
    recalc_section(db, section)
    db.flush()
    get_or_refresh_forming(db, section.id)
    get_or_refresh_labor(db, section.id)
    get_or_refresh_equipment(db, section.id)
    db.flush()
    recalc_section(db, section)
    db.flush()
    refresh_estimate_totals(db, estimate)
    db.flush()
    return section


def _summary(db, section) -> dict:
    chart, mapping = load_chart(db)
    return section_summary(db, section, chart, mapping)


def _lines(db, model, section_id) -> dict:
    return {r.code: r for r in db.scalars(select(model).where(model.section_id == section_id)).all() if r.enabled}


def _ext(rows: dict, *codes: str) -> Decimal:
    return sum((D(str(rows[c].ext_cost)) for c in codes if c in rows), D("0"))


# ------------------------------------------------------------- the matrix ----


@pytest.mark.parametrize("mod", SHAPES, ids=[m.replace("_fixture", "") for m in SHAPES])
def test_every_line_every_shape_writes_has_a_home_and_the_section_ties(db, estimate, mod):
    section = _full(db, estimate, mod)
    s = _summary(db, section)
    assert s["missing_line_sets"] == []
    assert s["unassigned"] == [], [(ln["kind"], ln["code"], ln["label"]) for ln in s["unassigned"]]
    assert s["cost"] > 0 and s["sale"] > 0
    # The per-row rounding, and only that: cents.
    assert abs(s["difference"]) < D("0.50"), (s["difference"], s["cost"])
    assert sum(s["columns"].values(), D("0")) + s["difference"] == s["cost"]
    counted = sum((v for c, v in s["codes"].items() if c not in SUBTOTALS and c != "000091"), D("0"))
    assert counted + s["tax"] + s["columns"]["unassigned"] + s["difference"] == s["cost"]
    assert s["profit"] + s["contingency"] + s["cost"] == s["sale"]
    assert s["codes"]["000040"] == sum((s["codes"][f"0000{n}"] for n in range(41, 50)), D("0"))
    assert s["codes"]["000060"] == sum((s["codes"][f"0000{n}"] for n in range(61, 69)), D("0"))
    assert set(s["columns"]) == set(COLUMN_KEYS)


# ---------------------------------------------------------- the columns ----


@pytest.fixture
def slab(db, estimate):
    return _full(db, estimate, "mono_slab_fixture")


def test_the_slab_columns_are_the_purchases_and_the_lines(db, slab):
    s = _summary(db, slab)
    mc = section_material_costs(db, slab)
    purchases = {ln["key"]: D(str(ln["cost"])) for ln in mc["lines"]}
    rounding = D(str(mc["rounding"]))
    assert s["columns"]["concrete"] == purchases["concrete"] + rounding   # the cents ride with the concrete
    assert s["columns"]["rebar"] == purchases["rebar"]
    assert s["columns"]["pt_cables"] == purchases["pt"]
    assert s["columns"]["grade_material"] == purchases["sand"]
    assert s["columns"]["poly_sealing"] == purchases["poly"] + purchases.get("tape", D("0"))
    assert s["columns"]["tax"] == D(str(slab.calc_total_tax)) > 0

    labor = _lines(db, EstimateLaborLine, slab.id)
    assert s["columns"]["supervision"] == _ext(labor, "superintendent", "foreman", "expense") > 0
    assert s["columns"]["pm"] == _ext(labor, "pm") > 0
    field = sum((D(str(r.ext_cost)) for r in labor.values() if r.group_name == "labor"), D("0"))
    assert slab.labor_subcontracted is False
    assert s["columns"]["labor"] == field > 0 and s["columns"]["sub_labor"] == 0
    assert s["codes"]["000020"] == _ext(labor, "superintendent", "foreman")
    assert s["codes"]["000021"] == _ext(labor, "expense") and s["codes"]["000022"] == _ext(labor, "pm")
    assert s["codes"]["000024"] == _ext(labor, "place_finish") > 0          # in house: the 20s block
    assert s["codes"]["000023"] == _ext(labor, "grading", "tie_steel", "excavation") > 0
    # No in-house forming code: the forming family stays in the 40s block either way.
    assert s["codes"]["000041"] == _ext(labor, "forming", "hold_downs", "brick_ledge", "thick_edge") > 0

    equipment = _lines(db, EstimateEquipmentLine, slab.id)
    assert s["columns"]["out_of_town"] == _ext(equipment, "out_of_town")
    assert s["codes"]["000030"] == _ext(equipment, "concrete_pump") > 0
    assert s["columns"]["other_subs"] >= _ext(equipment, "concrete_pump")
    fuel = sum((D(str(u.row.calc_equip_fuel or 0)) for u in cost_units(db, slab)), D("0"))
    assert s["codes"]["000066"] == fuel > 0
    rentals = sum((D(str(r.ext_cost)) for r in equipment.values() if (r.group_name or "equipment") == "equipment"), D("0"))
    assert s["columns"]["equipment"] == rentals + fuel + _ext(equipment, "easy_drill") * 0

    forming = _lines(db, EstimateFormingLine, slab.id)
    accessories = {"accessories", "chairs", "bolsters", "tie_wire", "dowels", "dowel_baskets", "smooth_dowels"}
    special = {"cure": "000010", "haul_off": "000033", "mesh": "000005", "rock": "000004", "carton_forms": "000009",
               "patch": "000011", "form_rental": "000039", "shoring_rental": "000039", "reshoring": "000039",
               "stud_rails": "000005", "bond_breaker": "000013", "french_drain": "000013",
               "lift_inserts": "000013", "brace_inserts": "000013"}
    lumber = sum((D(str(r.ext_cost)) for c, r in forming.items() if c not in accessories and c not in special), D("0"))
    assert s["codes"]["000002"] == lumber > 0
    assert s["codes"]["000006"] == sum((D(str(r.ext_cost)) for c, r in forming.items() if c in accessories), D("0")) > 0
    assert s["codes"]["000010"] == _ext(forming, "cure") + _ext(equipment, "cure")


def test_in_house_and_sub_labor_swap_columns_not_money(db, slab):
    before = _summary(db, slab)
    slab.labor_subcontracted = True
    db.flush()
    after = _summary(db, slab)
    assert before["columns"]["labor"] > 0 and before["columns"]["sub_labor"] == 0
    assert after["columns"]["sub_labor"] == before["columns"]["labor"] and after["columns"]["labor"] == 0
    assert after["cost"] == before["cost"] and after["sale"] == before["sale"]
    # The 20s block empties into the 40s: grading and tie steel to their sub codes.
    in_house = sum((before["codes"][f"0000{n}"] for n in range(23, 28)), D("0"))
    assert in_house > 0 and all(after["codes"][f"0000{n}"] == 0 for n in range(23, 28))
    assert after["codes"]["000040"] == before["codes"]["000040"] + in_house
    assert after["codes"]["000043"] == before["codes"]["000024"]
    assert after["codes"]["000042"] + after["codes"]["000047"] == before["codes"]["000023"]
    # Supervision is never subbed.
    for key in ("supervision", "pm", "concrete", "equipment", "other_subs"):
        assert after["columns"][key] == before["columns"][key]


# --------------------------------------------------------------- the job ----


@pytest.fixture
def job(db, estimate):
    """Three shapes on one job: the LBJ slab, the piers, an exercise of the misc library."""
    sections = [_full(db, estimate, m) for m in ("mono_slab_fixture", "piers_fixture", "misc_fixture")]
    for i, s in enumerate(sections):
        s.sort_order = (i + 1) * 10
    db.flush()
    refresh_estimate_totals(db, estimate)
    db.flush()
    return sections


def test_the_job_adds_its_sections_and_the_lower_block_is_the_tabs(client, db, project, estimate, job):
    r = client.get(f"/api/estimates/{estimate.id}/summary")
    assert r.status_code == 200, r.text
    s = r.json()
    assert [x["name"] for x in s["sections"]] == [x.name for x in job]
    lower = {k: D(str(v)) for k, v in s["lower"].items()}
    assert lower["contract_price"] == D(str(estimate.calc_total_sale)) > 0
    assert lower["cost"] == D(str(estimate.calc_total_cost)) > 0
    cols = {k: D(str(v)) for k, v in s["totals"]["columns"].items()}
    for k in COLUMN_KEYS:
        assert cols[k] == sum((D(str(x["columns"][k])) for x in s["sections"]), D("0"))
    assert lower["total_material"] == sum((cols[k] for k in MATERIAL_COLUMNS), D("0"))
    assert lower["total_labor"] == cols["labor"]
    assert lower["total_sub_labor"] == sum((cols[k] for k in SUB_LABOR_COLUMNS), D("0"))
    assert lower["total_other"] == sum((cols[k] for k in OTHER_COLUMNS), D("0"))
    assert lower["sales_tax"] == cols["tax"] and lower["unassigned"] == cols["unassigned"] == 0
    assert (lower["total_material"] + lower["total_labor"] + lower["total_sub_labor"] + lower["total_other"]
            + lower["sales_tax"] + lower["unassigned"] + lower["difference"]) == lower["cost"]
    assert lower["estimated_profit"] == lower["contract_price"] - lower["cost"] - lower["margin_contingency"]
    assert lower["labor_insurance"] == ((cols["sub_labor"] + cols["labor"] + cols["supervision"]) * LABOR_INSURANCE_PCT).quantize(D("0.01"))
    assert D(str(s["shares"]["concrete"])) == (cols["concrete"] / lower["contract_price"]).quantize(D("0.0001"))

    # 000091 is the contingency's share of the markup, on every section (the misc one's sale is typed).
    for x, section in zip(s["sections"], job):
        sale, cost = D(str(section.calc_total_sale)), D(str(section.calc_total_cost))
        m, c = D(str(section.margin_pct)), D(str(section.contingency_pct))
        expected = ((sale - cost) * c / (m + c)).quantize(D("0.01")) if m + c > 0 else D("0")
        assert D(str(x["codes"]["000091"])) == D(str(x["contingency"])) == expected
        assert D(str(x["profit"])) == sale - cost - expected
    slab = s["sections"][0]
    assert D(str(slab["contingency"])) == (D(str(job[0].calc_total_cost)) * D(str(job[0].contingency_pct))).quantize(D("0.01"))
    assert D(str(s["totals"]["codes"]["000091"])) == lower["margin_contingency"]

    # The chart, in order, and the header.
    assert [c["code"] for c in s["codes"]][:3] == ["000001", "000002", "000003"]
    assert s["codes"][-1]["code"] == "000091" and len(s["codes"]) == 58
    assert s["project_name"] == project.name and s["estimate_name"] == estimate.name
    assert s["file_name"].startswith(f"{project.name} - Summary - ") and s["file_name"].endswith(".xlsx")
    assert [c["key"] for c in s["columns"]] == list(COLUMN_KEYS)


# --------------------------------------------------------------- refiling ----


def test_a_senior_refiles_a_line_and_every_summary_follows(client, as_role, db, estimate, slab):
    pump = _ext(_lines(db, EstimateEquipmentLine, slab.id), "concrete_pump")
    assert pump > 0
    before = client.get(f"/api/estimates/{estimate.id}/summary").json()
    assert D(str(before["totals"]["codes"]["000030"])) == pump

    senior = as_role("senior_estimator")
    r = senior.put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000071"})
    assert r.status_code == 200, r.text
    assert r.json()["cost_code"] == "000071" and r.json()["seen"] is True and r.json()["label"] == "CONCRETE PUMPING"
    after = client.get(f"/api/estimates/{estimate.id}/summary").json()
    assert D(str(after["totals"]["codes"]["000030"])) == 0
    assert D(str(after["totals"]["codes"]["000071"])) == D(str(before["totals"]["codes"]["000071"])) + pump
    assert after["totals"]["columns"]["other_subs"] == before["totals"]["columns"]["other_subs"]  # both are other subs

    # Onto a rental code: the column moves with it, the cost does not.
    assert senior.put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000061"}).status_code == 200
    moved = client.get(f"/api/estimates/{estimate.id}/summary").json()
    assert D(str(moved["totals"]["columns"]["equipment"])) == D(str(before["totals"]["columns"]["equipment"])) + pump
    assert D(str(moved["totals"]["columns"]["other_subs"])) == D(str(before["totals"]["columns"]["other_subs"])) - pump
    assert moved["lower"]["cost"] == before["lower"]["cost"]

    # The list says where it is now.
    listed = client.get("/api/cost-codes").json()
    assert len(listed["codes"]) == 58
    line = next(ln for ln in listed["lines"] if ln["kind"] == "equipment" and ln["code"] == "concrete_pump")
    assert line["cost_code"] == "000061" and line["seen"] is True
    assert any(ln["kind"] == "purchase" and ln["code"] == "concrete" for ln in listed["lines"])

    # Refusals: the role, a code that is not one, a subtotal, an in-house code on a machine, a typo.
    assert as_role("estimator").put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000030"}).status_code == 403
    assert as_role("user").put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000030"}).status_code == 403
    assert as_role("user").get("/api/cost-codes").status_code == 200
    assert senior.put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000999"}).status_code == 400
    assert senior.put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000040"}).status_code == 400
    assert senior.put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_code": "000030", "inhouse_code": "000024"}).status_code == 400
    assert senior.put("/api/cost-codes/lines/gadget/concrete_pump", json={"cost_code": "000030"}).status_code == 404
    assert senior.put("/api/cost-codes/lines/equipment/concrete_pump", json={"cost_cod": "000030"}).status_code == 422
    # A labor line takes an in-house code, and a blank one clears it.
    r = senior.put("/api/cost-codes/lines/labor/forming", json={"cost_code": "000041", "inhouse_code": "000024"})
    assert r.status_code == 200 and r.json()["inhouse_code"] == "000024"
    r = senior.put("/api/cost-codes/lines/labor/forming", json={"cost_code": "000041", "inhouse_code": ""})
    assert r.status_code == 200 and r.json()["inhouse_code"] is None


def test_an_unknown_line_is_counted_as_unassigned_and_named(client, db, estimate, slab):
    db.add(EstimateLaborLine(
        section_id=slab.id, group_name="labor", code="bogus", label="SOMETHING NEW", enabled=True,
        rate=D("1"), unit="LS", qty=D("123.45"), ext_cost=D("123.45"), sort_order=999,
    ))
    db.flush()
    refresh_pour_costs(db, slab)
    refresh_estimate_totals(db, estimate)
    db.flush()
    s = _summary(db, slab)
    assert s["columns"]["unassigned"] == D("123.45")
    assert [(ln["kind"], ln["code"], ln["label"]) for ln in s["unassigned"]] == [("labor", "bogus", "SOMETHING NEW")]
    assert abs(s["difference"]) < D("0.50")
    assert sum(s["columns"].values(), D("0")) + s["difference"] == s["cost"]
    r = client.get(f"/api/estimates/{estimate.id}/summary").json()
    assert D(str(r["lower"]["unassigned"])) == D("123.45")
    listed = client.get("/api/cost-codes").json()
    line = next(ln for ln in listed["lines"] if ln["kind"] == "labor" and ln["code"] == "bogus")
    assert line["cost_code"] is None and line["seen"] is True and line["label"] == "SOMETHING NEW"


def test_a_section_nobody_opened_says_so(db, estimate):
    section = mf.build(db, estimate)
    refresh_pour_costs(db, section)
    refresh_estimate_totals(db, estimate)
    db.flush()
    s = _summary(db, section)
    assert s["missing_line_sets"] == ["forming", "labor", "equipment"]
    assert s["columns"]["concrete"] > 0 and s["columns"]["labor"] == 0 and s["columns"]["equipment"] == 0
    assert abs(s["difference"]) < D("0.50")


# ------------------------------------------------------------------ xlsx ----


def test_the_xlsx_is_the_tab(client, db, project, estimate, job):
    slab, piers, misc = job
    s = client.get(f"/api/estimates/{estimate.id}/summary").json()
    r = client.get(f"/api/estimates/{estimate.id}/summary.xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert s["file_name"] in r.headers["content-disposition"]

    ws = load_workbook(BytesIO(r.content))["Summary"]
    assert ws["A1"].value == "JOB COST SUMMARY"
    labels = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, 9)}
    assert labels["JOB NAME:"] == project.name

    # The job-cost table: a row per section with its sale and quantity, a $/unit line, the totals.
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    for section in job:
        r = rows[section.name]
        assert ws.cell(r, 4).value == float(section.calc_total_sale)
        assert ws.cell(r, 5).value == f"=IF(B{r}=0,0,D{r}/B{r})"
        assert ws.cell(r, 3).value == section.unit
        assert (ws.cell(r + 1, 1).value or "").startswith("$ / ")
    header = {ws.cell(rows[slab.name] - 1, c).value: c for c in range(1, ws.max_column + 1)}
    assert header["CONCRETE"] and header["SALES TAX"] and "UNASSIGNED" not in header
    assert ws.cell(rows[slab.name], header["CONCRETE"]).value == float(D(str(s["sections"][0]["columns"]["concrete"])))
    tot = rows["TOTALS"]
    assert ws.cell(tot, 4).value == "=" + "+".join(f"D{rows[x.name]}" for x in job)
    assert ws.cell(tot, header["CONCRETE"]).value.startswith("=")

    # The lower block, in the tab's order, and the profit as the price less the rest.
    lower = {ws.cell(r, 2).value: r for r in range(tot, ws.max_row + 1) if ws.cell(r, 2).value}
    for name in ("TOTAL MATERIAL", "TOTAL LABOR", "TOTAL SUB LABOR", "TOTAL OTHER", "SALES TAX",
                 "MARGIN CONTINGENCY", "ESTIMATED PROFIT", "CONTRACT PRICE"):
        assert name in lower, name
    assert ws.cell(lower["MARGIN CONTINGENCY"], 4).value == float(D(str(s["lower"]["margin_contingency"])))
    assert ws.cell(lower["ESTIMATED PROFIT"], 4).value.startswith(f"=D{tot}-SUM(D")
    assert ws.cell(lower["CONTRACT PRICE"], 4).value == f"=D{tot}"

    # The cost codes: the chart in order, a column per section, the tie-out under it.
    codes = [ws.cell(r, 2).value for r in range(lower["CONTRACT PRICE"], ws.max_row + 1)
             if isinstance(ws.cell(r, 2).value, str) and ws.cell(r, 2).value.startswith("0000")]
    assert codes == [c["code"] for c in s["codes"]]
    code_rows = {ws.cell(r, 2).value: r for r in range(lower["CONTRACT PRICE"], ws.max_row + 1)}
    r91 = code_rows["000091"]
    assert ws.cell(r91, 5).value == float(D(str(s["sections"][0]["contingency"])))
    assert ws.cell(code_rows["000008"], 5).value == float(D(str(s["sections"][0]["codes"]["000008"])))
    assert ws.cell(code_rows["000040"], 5).value.startswith("=E")      # a subtotal is a formula over its block
    assert ws.cell(code_rows["000008"], 4).value == f"=SUM(E{code_rows['000008']}:G{code_rows['000008']})"
    names = {ws.cell(r, 3).value: r for r in range(r91, ws.max_row + 1)}
    assert ws.cell(names["SALE"], 5).value == float(slab.calc_total_sale)
    assert ws.cell(names["ESTIMATED PROFIT"], 5).value == f"=E{names['SALE']}-E{names['TOTAL COST']}-E{r91}"
