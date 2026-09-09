"""
The proposal (sql/080): the estimate pushed onto the bid form.

Chad, 2026-09-08: "ok, proposal? here is the latest proposal form we are
using" — Sherry Pointe Apts 26-102 on the NEW FORM — and, on the proposal,
"build it".

What is pinned here:

  * a proposal SEEDS from the takeoff — one section per estimate section in
    order, one line per costed row across the shapes, the row's name, its
    quantity in the section's unit, its sale per unit — and the tie-out
    beside each section holds to the rounding of a unit price to four places;
  * a REFRESH moves the numbers and keeps the words: a rewritten description
    and an excluded status survive, a deleted row flags its line, a new row
    and a new section arrive;
  * the grids, the blocks and the header edit; a typo is a 422; a blank on a
    field that cannot be empty leaves it;
  * the .xlsx IS the form — the letterhead, the bands, the header row, the
    numbered lines with `=C*F`, `=SUM` under each section, the lump sum over
    the section totals, the blocks and the terms in order, the print area
    stopping at E so the prices never print;
  * the library is the company's text: a senior estimator edits it, a new
    proposal copies it, a proposal already made keeps its own.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.estimate import Estimate
from app.models.estimate_section import EstimateSection
from app.models.mono_slab import MonoSlab
from app.services.costing import refresh_estimate_totals, refresh_pour_costs
from tests import misc_fixture as mcf
from tests import mono_slab_fixture as mf
from tests import piers_fixture as pif

D = Decimal
LIBRARY_COUNTS = {
    "alternates": 8, "equipment_rates": 9, "labor_rates": 3,
    "qualifications": 8, "exclusions": 8, "terms": 6,
}


@pytest.fixture
def job(db, estimate):
    """Three shapes on one job: the LBJ slab, the piers, an exercise of the misc library."""
    sections = [mf.build(db, estimate), pif.build(db, estimate), mcf.build(db, estimate)]
    # Built in one transaction they share a created_at, so the order is pinned
    # the way a job's sections are — by sort_order.
    for i, s in enumerate(sections):
        s.sort_order = (i + 1) * 10
        refresh_pour_costs(db, s)
    refresh_estimate_totals(db, estimate)
    db.flush()
    return sections


def _make(client, estimate) -> dict:
    r = client.post("/api/proposals", json={"estimate_id": str(estimate.id)})
    assert r.status_code == 201, r.text
    return r.json()


def _tolerance(section: dict) -> Decimal:
    """A unit price to four places moves a line by up to half a ten-thousandth per unit."""
    lines = section["lines"]
    return sum((D(str(ln["qty"])) for ln in lines), D("0")) * D("0.00005") + D("0.01") * len(lines)


# ---------------------------------------------------------------- seed ----


def test_a_proposal_seeds_from_the_takeoff(client, db, project, estimate, job):
    slab, piers, misc = job
    p = _make(client, estimate)

    assert [s["title"] for s in p["sections"]] == [slab.name, piers.name, misc.name]
    assert [s["section_id"] for s in p["sections"]] == [str(slab.id), str(piers.id), str(misc.id)]

    # One line per pour, in order, carrying the pour's own numbers.
    mono = p["sections"][0]
    pours = list(db.scalars(
        select(MonoSlab).where(MonoSlab.section_id == slab.id).order_by(MonoSlab.sort_order, MonoSlab.created_at)
    ).all())
    assert len(mono["lines"]) == len(pours) == len(mf.POURS)
    for line, pour in zip(mono["lines"], pours):
        # The seeded description says what the thing is: the pour's name,
        # then its thickness and its PT in the form's idiom.
        assert line["description"].startswith(pour.description + ' (4" PT SOG')
        assert line["source_label"] == pour.description
        assert D(str(line["qty"])) == D(str(pour.square_footage)) * pour.qty
        assert line["unit"] == "SF" and line["status"] == "INCLUDED"
        assert (line["source_table"], line["source_id"]) == ("mono_slabs", str(pour.id))
        assert D(str(line["unit_price"])) == (D(str(pour.calc_sale)) / D(str(pour.square_footage))).quantize(D("0.0001"))
        assert D(str(line["extended"])) == (D(str(line["qty"])) * D(str(line["unit_price"]))).quantize(D("0.01"))

    # The tie-out: each section against its estimate section, the whole
    # against the job, within the rounding of a four-place unit price.
    for s in p["sections"]:
        assert s["estimate_sale"] is not None and D(str(s["estimate_sale"])) > 0
        assert abs(D(str(s["difference"]))) <= _tolerance(s), (s["title"], s["difference"])
    assert D(str(p["estimate_sale"])) == D(str(estimate.calc_total_sale)).quantize(D("0.01"))
    assert abs(D(str(p["difference"]))) <= sum((_tolerance(s) for s in p["sections"]), D("0"))
    assert D(str(p["total"])) == sum((D(str(s["total"])) for s in p["sections"]), D("0"))

    # A library item nobody used is not a line; a used one is, in its own unit.
    used = [code for code, qty in mcf.QUANTITIES.items() if D(str(qty)) > 0]
    assert len(p["sections"][2]["lines"]) == len(used)
    assert all(ln["unit"] == misc.unit for ln in p["sections"][2]["lines"])
    # A pier group and a round base describe themselves from their own fields.
    pier_line = p["sections"][1]["lines"][0]
    assert '" dia x ' in pier_line["description"] and "' deep" in pier_line["description"]
    assert any('" dia x ' in ln["description"] for ln in p["sections"][2]["lines"])

    # The header from the project, the four drawings, the playbook's file name.
    assert (p["submitted_to"], p["job_label"], p["location"]) == (project.gc, project.name, project.location)
    assert [d["discipline"] for d in p["drawings"]] == ["ARCHITECTURAL", "STRUCTURAL", "CIVIL", "LANDSCAPE"]
    assert (p["rev"], p["proposal_date"]) == (1, date.today().isoformat())
    assert p["file_name"] == f"{project.name} - Proposal - {date.today():%Y-%m-%d}_01.xlsx"
    assert p["intro"].startswith("We are pleased to submit") and p["payment_terms"] == "PER CONTRACT AGREEMENT OR UPON COMPLETION"

    # The standing text, copied from the library.
    assert {b: len(v) for b, v in p["items"].items()} == LIBRARY_COUNTS

    # One per estimate, found by the estimate.
    assert client.post("/api/proposals", json={"estimate_id": str(estimate.id)}).status_code == 409
    assert client.get(f"/api/proposals?estimate_id={estimate.id}").json()["id"] == p["id"]


def test_a_job_number_is_in_the_name(client, db, estimate, job):
    from app.models.project import Project

    project = db.get(Project, estimate.project_id)
    project.job_number = "26-102"
    project.gc = "Spring Valley Construction Company"
    db.flush()
    p = _make(client, estimate)
    assert p["job_label"] == f"{project.name}  ·  26-102"
    assert p["submitted_to"] == "Spring Valley Construction Company"
    assert p["file_name"] == f"{project.name} - Proposal - 26-102 - {date.today():%Y-%m-%d}_01.xlsx"


# ------------------------------------------------------------- refresh ----


def test_a_refresh_moves_the_numbers_and_keeps_the_words(client, db, estimate, job):
    slab = job[0]
    p = _make(client, estimate)
    mono = p["sections"][0]
    first, second, third = mono["lines"][0], mono["lines"][1], mono["lines"][2]

    # The estimator rewrites one description and excludes another line.
    r = client.put(f"/api/proposal-sections/{mono['id']}/lines/bulk", json={"rows": [
        {"id": first["id"], "description": 'Building A (4" PT SOG 3500 psi w/ 12" x 36" perimeter grade beams)'},
        {"id": second["id"], "status": "EXCLUDED"},
    ]})
    assert r.status_code == 200, r.text

    # The takeoff moves: a thousand square feet onto the first pour.
    r = client.patch(
        f"/api/mono-slabs/{first['source_id']}",
        json={"square_footage": float(D(str(first["qty"])) + 1000)},
    )
    assert r.status_code == 200, r.text
    res = client.post(f"/api/proposals/{p['id']}/refresh").json()
    assert res["updated"] >= 1 and res["added_lines"] == 0 and res["added_sections"] == 0 and res["missing"] == 0
    lines = {ln["id"]: ln for ln in res["proposal"]["sections"][0]["lines"]}
    moved = lines[first["id"]]
    pour = db.get(MonoSlab, UUID(first["source_id"]))
    assert D(str(moved["qty"])) == D(str(first["qty"])) + 1000 == D(str(pour.square_footage))
    assert D(str(moved["unit_price"])) == (D(str(pour.calc_sale)) / D(str(pour.square_footage))).quantize(D("0.0001"))
    assert moved["description"].startswith("Building A (4")          # the words stayed
    assert lines[second["id"]]["status"] == "EXCLUDED"                # so did the status
    assert D(str(lines[second["id"]]["extended"])) == 0 and lines[second["id"]]["qty"] is not None

    # A pour deleted flags its line; a pour added and a section added arrive.
    assert client.delete(f"/api/mono-slabs/{third['source_id']}").status_code == 204
    r = client.post("/api/mono-slabs", json={
        "section_id": str(slab.id), "description": "Pour 26", "square_footage": 1234, "thickness_in": 4,
    })
    assert r.status_code == 201, r.text
    courtyard = EstimateSection(estimate_id=estimate.id, kind="mono_slab", name="Pool courtyard", unit="SF")
    db.add(courtyard)
    db.flush()
    res = client.post(f"/api/proposals/{p['id']}/refresh").json()
    assert (res["missing"], res["added_lines"], res["added_sections"]) == (1, 1, 1)
    mono2 = res["proposal"]["sections"][0]
    flagged = next(ln for ln in mono2["lines"] if ln["id"] == third["id"])
    assert flagged["source_missing"] is True and flagged["source_label"] is None
    assert any(ln["description"].startswith('Pour 26 (4" SOG') and D(str(ln["qty"])) == 1234 for ln in mono2["lines"])
    assert res["proposal"]["sections"][-1]["title"] == "Pool courtyard"
    assert res["proposal"]["sections"][-1]["lines"] == []
    # A second refresh changes nothing.
    res = client.post(f"/api/proposals/{p['id']}/refresh").json()
    assert (res["updated"], res["added_lines"], res["added_sections"], res["missing"]) == (0, 0, 0, 1)


# ------------------------------------------------------ grids and header ----


def test_the_grids_the_blocks_and_the_header(client, db, project, estimate, job):
    p = _make(client, estimate)
    pid = p["id"]
    first = p["sections"][0]

    # Sections: the first renamed, one added by hand, the rest dropped.
    r = client.put(f"/api/proposals/{pid}/sections/bulk", json={"rows": [
        {"id": first["id"], "title": "BUILDING FOUNDATIONS - POST-TENSIONED SLAB ON GRADE"},
        {"title": "MEASURED AND EXCLUDED - SCOPE BY OTHERS"},
    ], "delete_missing": True})
    assert r.status_code == 200, r.text
    got = r.json()
    assert [s["title"] for s in got["sections"]] == [
        "BUILDING FOUNDATIONS - POST-TENSIONED SLAB ON GRADE", "MEASURED AND EXCLUDED - SCOPE BY OTHERS",
    ]
    hand = got["sections"][1]
    assert hand["section_id"] is None and hand["estimate_sale"] is None and hand["lines"] == []
    assert got["sections"][0]["lines"] == first["lines"]   # renaming a section touches no line

    # Lines typed by hand: one priced, one measured and excluded with its quantity kept.
    r = client.put(f"/api/proposal-sections/{hand['id']}/lines/bulk", json={"rows": [
        {"description": "Spoils Haul Off (76 truck loads, 1,500 YDS)", "qty": 1, "unit": "LS",
         "status": "INCLUDED", "unit_price": 35000},
        {"description": "Lime Stabilized Subgrade", "qty": 19422.86, "unit": "SF", "status": "EXCLUDED"},
    ]})
    assert r.status_code == 200, r.text
    sec = r.json()["sections"][1]
    assert [D(str(ln["extended"])) for ln in sec["lines"]] == [D("35000.00"), D("0.00")]
    assert D(str(sec["total"])) == D("35000.00") and sec["difference"] is None
    assert D(str(sec["lines"][1]["qty"])) == D("19422.86") and sec["lines"][1]["unit_price"] is None
    assert all(ln["source_id"] is None for ln in sec["lines"])
    assert D(str(r.json()["total"])) == D(str(got["sections"][0]["total"])) + D("35000.00")

    # A block, whole; a blank string is not a bullet; an unknown block is a 404.
    r = client.put(f"/api/proposals/{pid}/items/labor_rates",
                   json={"items": ["Laborers - $38.00 / hr", "  ", "Supervision - $210.00 / hr"]})
    assert r.status_code == 200, r.text
    assert [i["text"] for i in r.json()["items"]["labor_rates"]] == ["Laborers - $38.00 / hr", "Supervision - $210.00 / hr"]
    assert client.put(f"/api/proposals/{pid}/items/nope", json={"items": []}).status_code == 404

    # The header, the drawings, the revision and its file name.
    r = client.patch(f"/api/proposals/{pid}", json={
        "attn": "Jamie Daigneault, Estimator", "email": "jdaigneault@svcc.biz", "rev": 2,
        "proposal_date": "2026-09-01",
        "drawings": [
            {"discipline": "ARCHITECTURAL", "firm": "Cross Architects, PLLC", "plan_date": "2025-06-05"},
            {"discipline": "STRUCTURAL", "firm": "", "plan_date": ""},
        ],
    })
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["rev"] == 2 and got["file_name"] == f"{project.name} - Proposal - 2026-09-01_02.xlsx"
    assert got["drawings"][0]["plan_date"] == "2025-06-05" and got["drawings"][1]["firm"] is None
    # A blank on a field that cannot be empty leaves it alone.
    r = client.patch(f"/api/proposals/{pid}", json={"rev": "", "intro": "", "attn": ""})
    assert r.status_code == 200, r.text
    assert r.json()["rev"] == 2 and r.json()["intro"] == got["intro"] and r.json()["attn"] == ""

    # A misspelled field is a 422 on every write model.
    for method, url, body in [
        ("POST", "/api/proposals", {"estimate_id": str(estimate.id), "estimat_id": 1}),
        ("PATCH", f"/api/proposals/{pid}", {"atn": "x"}),
        ("PUT", f"/api/proposals/{pid}/sections/bulk", {"rows": [{"titel": "x"}]}),
        ("PUT", f"/api/proposal-sections/{hand['id']}/lines/bulk", {"rows": [{"descripton": "x"}]}),
        ("PUT", f"/api/proposal-sections/{hand['id']}/lines/bulk", {"rows": [{"status": "MAYBE"}]}),
        ("PUT", f"/api/proposals/{pid}/items/terms", {"itmes": []}),
        ("PUT", "/api/proposal-library/terms", {"itmes": []}),
    ]:
        assert client.request(method, url, json=body).status_code == 422, (method, url)

    # Deletes, down to the proposal itself.
    assert client.delete(f"/api/proposal-lines/{sec['lines'][0]['id']}").status_code == 204
    assert len(client.get(f"/api/proposals/{pid}").json()["sections"][1]["lines"]) == 1
    assert client.delete(f"/api/proposal-sections/{hand['id']}").status_code == 204
    assert len(client.get(f"/api/proposals/{pid}").json()["sections"]) == 1
    assert client.delete(f"/api/proposals/{pid}").status_code == 204
    assert client.get(f"/api/proposals?estimate_id={estimate.id}").status_code == 404


# ---------------------------------------------------------------- xlsx ----


def test_the_xlsx_is_the_form(client, db, estimate, job):
    slab = job[0]
    p = _make(client, estimate)
    r = client.get(f"/api/proposals/{p['id']}/xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert p["file_name"] in r.headers["content-disposition"]

    ws = load_workbook(BytesIO(r.content))["PROPOSAL"]
    cells = {(c.row, c.column_letter): c.value for row in ws.iter_rows() for c in row if c.value is not None}

    # The letterhead, the widths, the page.
    assert cells[(1, "A")] == "S & S CONCRETE CONTRACTORS, INC." and cells[(6, "B")] == "PROPOSAL"
    assert ws.column_dimensions["B"].width == 74 and ws.column_dimensions["G"].width == 15
    assert ws.oddFooter.right.text == "CONFIDENTIAL" and ws.oddFooter.center.text == "Page &P of &N"
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.orientation == "portrait"
    # The print area stops at E: unit prices and extendeds never print.
    last_col, last_row = ws.print_area.split(":")[-1].lstrip("$").split("$")
    assert last_col == "E" and int(last_row) >= ws.max_row

    # The header block and the drawings.
    b = [v for (row, col), v in sorted(cells.items()) if col == "B" and isinstance(v, str)]
    assert any(v.startswith("DATE     " + date.today().isoformat()) for v in b)
    assert any(v.startswith("JOB     ") for v in b) and "DRAWINGS SUPPLIED" in b
    assert any(v.startswith("ARCHITECTURAL:") for v in b) and any(v.startswith("LANDSCAPE:") for v in b)

    # The first section: its band, its header row, its numbered lines, its total.
    (title_row,) = [row for (row, col), v in cells.items() if col == "B" and v == slab.name.upper()]
    head = title_row + 1
    assert [cells[(head, c)] for c in "ABCDEFG"] == ["#", "DESCRIPTION", "QTY", "UNIT", "STATUS", "UNIT PRICE", "EXTENDED"]
    lines = p["sections"][0]["lines"]
    first, r1 = lines[0], head + 1
    assert cells[(r1, "A")] == 1 and cells[(r1, "B")] == first["description"]
    assert cells[(r1, "C")] == int(D(str(first["qty"]))) and cells[(r1, "D")] == "SF" and cells[(r1, "E")] == "INCLUDED"
    assert cells[(r1, "F")] == float(D(str(first["unit_price"]))) and cells[(r1, "G")] == f"=C{r1}*F{r1}"
    total_row = r1 + len(lines)
    assert cells[(total_row - 1, "A")] == len(lines)
    assert cells[(total_row, "D")] == "SECTION TOTAL"
    assert cells[(total_row, "E")] == f"=SUM(G{r1}:G{total_row - 1})"
    # Numbering runs on across sections, the way the form's does.
    numbers = [v for (row, col), v in sorted(cells.items()) if col == "A" and isinstance(v, int)]
    n_lines = sum(len(s["lines"]) for s in p["sections"])
    assert numbers[:n_lines] == list(range(1, n_lines + 1))

    # The lump sum adds the three section totals.
    lump = [v for (row, col), v in cells.items() if col == "E" and isinstance(v, str) and v.startswith("=E")]
    assert len(lump) == 1 and lump[0].count("+") == 2

    # The blocks and the acceptance page, in the form's order, with the library's text.
    order = ["ALTERNATES", "ADDITIONAL EQUIPMENT RATES", "ADDITIONAL LABOR RATES", "QUALIFICATIONS",
             "EXCLUSIONS", "TERMS AND CONDITIONS", "ACCEPTANCE", "Date of Acceptance", "Customer Signature",
             "Authorized Signature, S & S Concrete Contractors, Inc."]
    idx = [b.index(x) for x in order]
    assert idx == sorted(idx)
    assert "Laborers - $36.00 / hr" in b and "Any item not specifically included is excluded." in b
    assert any(v.startswith("General Contractor is responsible") for v in b)
    assert any(v.startswith("Payment to be made as follows:  PER CONTRACT") for v in b)
    bullets = [v for (row, col), v in cells.items() if col == "A" and v == "•"]
    assert len(bullets) == 8 + 9 + 3 + 8 + 8
    # A wrapped line is given its room: the force majeure paragraph is the tallest row.
    (fm_row,) = [row for (row, col), v in cells.items() if col == "B" and isinstance(v, str) and v.startswith("An event of")]
    assert ws.row_dimensions[fm_row].height > 150


def test_an_excluded_line_prints_its_quantity_and_no_price(client, db, estimate, job):
    p = _make(client, estimate)
    mono = p["sections"][0]
    line = mono["lines"][0]
    client.put(f"/api/proposal-sections/{mono['id']}/lines/bulk", json={"rows": [{"id": line["id"], "status": "EXCLUDED"}]})
    ws = load_workbook(BytesIO(client.get(f"/api/proposals/{p['id']}/xlsx").content))["PROPOSAL"]
    rows = [c.row for c in ws["B"] if c.value == line["description"]]
    assert len(rows) == 1
    r = rows[0]
    assert ws[f"E{r}"].value == "EXCLUDED" and ws[f"C{r}"].value == int(D(str(line["qty"])))
    assert ws[f"F{r}"].value is None and ws[f"G{r}"].value is None


# ------------------------------------------------------------- library ----


def test_the_library_is_the_companys_text(client, as_role, db, estimate, job):
    lib = client.get("/api/proposal-library").json()["blocks"]
    assert {b: len(v) for b, v in lib.items()} == LIBRARY_COUNTS
    assert lib["labor_rates"][0]["text"] == "Laborers - $36.00 / hr"
    assert lib["alternates"][0]["text"].startswith("Pipe Bollards")
    assert not any("Sales Tax Deduction" in i["text"] for i in lib["alternates"])

    p = _make(client, estimate)
    senior = as_role("senior_estimator")
    r = senior.put("/api/proposal-library/labor_rates",
                   json={"items": ["Laborers - $38.00 / hr", "Supervision - $210.00 / hr"]})
    assert r.status_code == 200, r.text
    assert [i["text"] for i in r.json()["blocks"]["labor_rates"]] == ["Laborers - $38.00 / hr", "Supervision - $210.00 / hr"]

    # The proposal already made keeps its copy; a new one copies the new text.
    assert len(client.get(f"/api/proposals/{p['id']}").json()["items"]["labor_rates"]) == 3
    second = Estimate(project_id=estimate.project_id, name="Second look")
    db.add(second)
    db.flush()
    p2 = _make(client, second)
    assert [i["text"] for i in p2["items"]["labor_rates"]] == ["Laborers - $38.00 / hr", "Supervision - $210.00 / hr"]
    assert p2["sections"] == [] and D(str(p2["total"])) == 0 and p2["estimate_sale"] is None

    # An estimator reads the library and makes proposals; the company's text is a senior's.
    estimator = as_role("estimator")
    assert estimator.get("/api/proposal-library").status_code == 200
    assert estimator.put("/api/proposal-library/labor_rates", json={"items": ["x"]}).status_code == 403
    assert estimator.post(f"/api/proposals/{p2['id']}/refresh").status_code == 200
    viewer = as_role("user")
    assert viewer.get(f"/api/proposals/{p2['id']}/xlsx").status_code == 200
    assert viewer.post(f"/api/proposals/{p2['id']}/refresh").status_code == 403


# ---------------------------------------------------- walls and footings ----


def test_a_wall_run_is_two_lines(client, db, estimate):
    """
    Chad, 2026-09-08: "I want the walls and associated footings separate."
    The wall on its form feet at the wall's sale, the footing on its length at
    the footing's — the split the costing keeps — and the two are the run.
    """
    from app.models.wall_run import WallRun
    from app.services.recalc import recalc_section
    from tests import walls_fixture as wf

    section = wf.build(db, estimate)
    refresh_pour_costs(db, section)
    refresh_estimate_totals(db, estimate)
    db.flush()
    p = _make(client, estimate)
    (walls,) = p["sections"]
    runs = list(db.scalars(
        select(WallRun).where(WallRun.section_id == section.id).order_by(WallRun.sort_order, WallRun.created_at)
    ).all())
    assert len(runs) == 16 and len(walls["lines"]) == 32
    for i, run in enumerate(runs):
        wall, ftg = walls["lines"][2 * i], walls["lines"][2 * i + 1]
        assert (wall["source_part"], ftg["source_part"]) == ("wall", "footing")
        assert wall["source_id"] == ftg["source_id"] == str(run.id)
        assert wall["description"].startswith(f"{run.label}: ") and '" tall x ' in wall["description"]
        assert ftg["description"].startswith(f"{run.label} footing: ") and '" footing' in ftg["description"]
        assert (wall["unit"], ftg["unit"]) == ("FF", "LF")
        assert D(str(wall["qty"])) == D(str(run.calc_form_ff)).quantize(D("0.001"))
        assert D(str(ftg["qty"])) == D(str(run.length_ft)).quantize(D("0.001"))
        assert D(str(wall["unit_price"])) == (D(str(run.calc_wall_sale)) / D(str(run.calc_form_ff))).quantize(D("0.0001"))
        assert D(str(ftg["unit_price"])) == (D(str(run.calc_footing_sale)) / D(str(run.length_ft))).quantize(D("0.0001"))
        assert wall["source_label"] == f"{run.label} · wall" and ftg["source_label"] == f"{run.label} · footing"
        # The two halves are the run, to the rounding of two four-place prices.
        slack = D("0.02") + (D(str(run.calc_form_ff)) + D(str(run.length_ft))) * D("0.00005")
        assert abs(D(str(wall["extended"])) + D(str(ftg["extended"])) - D(str(run.calc_sale))) <= slack, run.label
    assert abs(D(str(walls["difference"]))) <= _tolerance(walls)

    # A run stripped of its footing: the footing line is flagged on refresh,
    # the wall line moves with the split, nothing is added.
    first = runs[0]
    first.ftg_width_in = D("0")
    first.ftg_thick_in = D("0")
    recalc_section(db, section)
    db.flush()
    res = client.post(f"/api/proposals/{p['id']}/refresh").json()
    assert (res["missing"], res["added_lines"], res["added_sections"]) == (1, 0, 0)
    lines = res["proposal"]["sections"][0]["lines"]
    assert lines[1]["source_missing"] is True and lines[0]["source_missing"] is False
    assert D(str(lines[0]["unit_price"])) == (D(str(first.calc_wall_sale)) / D(str(first.calc_form_ff))).quantize(D("0.0001"))


def test_spot_footings_sell_per_each(client, db, estimate):
    """Sold per footing in the costing since sql/072; the label said SF until sql/081."""
    from app.models.wall_run import WallRun
    from tests import spot_footings_fixture as sff

    section = sff.build(db, estimate)
    refresh_pour_costs(db, section)
    refresh_estimate_totals(db, estimate)
    db.flush()
    p = _make(client, estimate)
    (spots,) = p["sections"]
    runs = list(db.scalars(select(WallRun).where(WallRun.section_id == section.id)).all())
    assert len(spots["lines"]) == len(runs) == 4
    for ln in spots["lines"]:
        assert ln["unit"] == "EA" and ln["source_part"] is None
        assert D(str(ln["qty"])) == D(str(ln["qty"])).to_integral_value()
    assert sum(D(str(ln["qty"])) for ln in spots["lines"]) == sum(D(str(r.footing_count)) for r in runs)
    # A new spot footings section defaults to EA, not SF.
    r = client.post(f"/api/estimates/{estimate.id}/sections", json={"kind": "spot_footings", "name": "Pads"})
    assert r.status_code == 201, r.text
    assert r.json()["unit"] == "EA"


def test_a_line_seeded_before_the_split_becomes_the_wall_half(client, db, estimate):
    """The live job's proposal predates sql/081: its wall lines carry no part. A refresh adopts each as the wall and adds the footing."""
    from app.models.proposal import ProposalLine
    from tests import walls_fixture as wf

    section = wf.build(db, estimate)
    refresh_pour_costs(db, section)
    refresh_estimate_totals(db, estimate)
    db.flush()
    p = _make(client, estimate)
    wall, ftg = p["sections"][0]["lines"][0], p["sections"][0]["lines"][1]
    # Undo the split on the first run, the way a pre-081 proposal stands.
    old = db.get(ProposalLine, UUID(wall["id"]))
    old.source_part = None
    old.description = "W1 as it was written before the split"
    db.delete(db.get(ProposalLine, UUID(ftg["id"])))
    db.flush()
    res = client.post(f"/api/proposals/{p['id']}/refresh").json()
    assert (res["missing"], res["added_lines"]) == (0, 1)
    lines = res["proposal"]["sections"][0]["lines"]
    adopted = next(ln for ln in lines if ln["id"] == wall["id"])
    assert adopted["source_part"] == "wall" and adopted["description"] == "W1 as it was written before the split"
    assert sum(1 for ln in lines if ln["source_id"] == wall["source_id"] and ln["source_part"] == "footing") == 1


# ------------------------------------------------------------- filing ----


def test_a_line_is_filed_where_the_work_is(client, db, estimate, job):
    """
    Chad, 2026-09-08: "is it possible to build the proposal where it is
    instead of just the section... like spot footings that are in the mono
    slab to be added to that section". A line goes where the estimator files
    it; the tie-out follows; a refresh remembers.
    """
    slab, piers, misc = job
    p = _make(client, estimate)
    mono, pier_sec, misc_sec = p["sections"]
    total_before = D(str(p["total"]))

    # One pier group filed under the slab: its words and numbers travel with it.
    moved = pier_sec["lines"][0]
    r = client.put(f"/api/proposal-sections/{pier_sec['id']}/lines/bulk", json={"rows": [
        {"id": moved["id"], "proposal_section_id": mono["id"], "description": "Elevator pit piers"},
    ]})
    assert r.status_code == 200, r.text
    got = r.json()
    mono2, pier2 = got["sections"][0], got["sections"][1]
    assert mono2["lines"][-1]["id"] == moved["id"] and mono2["lines"][-1]["proposal_section_id"] == mono["id"]
    assert mono2["lines"][-1]["description"] == "Elevator pit piers"
    assert D(str(mono2["lines"][-1]["unit_price"])) == D(str(moved["unit_price"]))
    assert all(ln["id"] != moved["id"] for ln in pier2["lines"])
    assert D(str(got["total"])) == total_before

    # The tie-out follows the line: the slab section's estimate figure grew by
    # the pier group's sale, the piers section's shrank by it.
    grew = D(str(mono2["estimate_sale"])) - D(str(mono["estimate_sale"]))
    shrank = D(str(pier_sec["estimate_sale"])) - D(str(pier2["estimate_sale"]))
    assert grew == shrank and abs(grew - D(str(moved["extended"]))) <= D("0.05")
    assert abs(D(str(mono2["difference"]))) <= _tolerance(mono2)

    # A typed line is money on top; a section of typed lines has no estimate figure.
    r = client.put(f"/api/proposals/{p['id']}/sections/bulk", json={"rows": [{"title": "SITE WORK"}]})
    site = r.json()["sections"][-1]
    r = client.put(f"/api/proposal-sections/{site['id']}/lines/bulk", json={"rows": [
        {"description": "Certified Payroll", "qty": 1, "unit": "LS", "unit_price": 7500},
    ]})
    site = r.json()["sections"][-1]
    assert site["estimate_sale"] is None and D(str(site["total"])) == D("7500.00")

    # Every misc line moved at once, the emptied section let go, and a refresh
    # files a new misc item beside its siblings rather than reviving the section.
    r = client.post(f"/api/proposal-sections/{misc_sec['id']}/move-lines", json={"to": site["id"]})
    assert r.status_code == 200, r.text
    site2 = next(s for s in r.json()["sections"] if s["id"] == site["id"])
    assert len(site2["lines"]) == 1 + len(misc_sec["lines"])
    assert [ln["id"] for ln in site2["lines"][1:]] == [ln["id"] for ln in misc_sec["lines"]]
    assert client.delete(f"/api/proposal-sections/{misc_sec['id']}").status_code == 204
    r = client.post("/api/misc-items", json={
        "section_id": str(misc.id), "code": "1399", "description": "Flagpole base", "shape": "round",
        "unit": "EA", "qty": 2, "unit_sale": 900, "labor_per_unit": 150, "dim_a": 24, "dim_b": 4,
    })
    assert r.status_code == 201, r.text
    res = client.post(f"/api/proposals/{p['id']}/refresh").json()
    assert (res["added_lines"], res["added_sections"]) == (1, 0)
    site3 = next(s for s in res["proposal"]["sections"] if s["id"] == site["id"])
    assert site3["lines"][-1]["description"].startswith("Flagpole base (")
    assert not any(s["title"] == misc.name for s in res["proposal"]["sections"])

    # A move needs a section of the same proposal, and not the same one.
    other = Estimate(project_id=estimate.project_id, name="Other job")
    db.add(other)
    db.flush()
    p2 = _make(client, other)
    r = client.put(f"/api/proposals/{p2['id']}/sections/bulk", json={"rows": [{"title": "ELSEWHERE"}]})
    elsewhere = r.json()["sections"][0]
    r = client.put(f"/api/proposal-sections/{mono['id']}/lines/bulk", json={"rows": [
        {"id": mono["lines"][0]["id"], "proposal_section_id": elsewhere["id"]},
    ]})
    assert r.status_code == 400
    assert client.post(f"/api/proposal-sections/{mono['id']}/move-lines", json={"to": mono["id"]}).status_code == 400
    assert client.post(f"/api/proposal-sections/{mono['id']}/move-lines", json={"to": elsewhere["id"]}).status_code == 400

