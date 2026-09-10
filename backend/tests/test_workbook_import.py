"""
An estimate workbook read onto a job (services/workbook_import.py).

Chad, 2026-09-10, with the Lakeside Townhomes workbook: "lets import this
estimate" — "full import, build it".

The workbook here is built by the test in the Lakeside template's layout —
the headers a row lower than the LBJ tab's, a QTY column on the slab tab,
the labor block's costs under a "SUB LABOR" header the tabs merge across two
cells — because the reader is keyed by header text and that is the contract
worth pinning. What is pinned:

  * the header, the rows (a pour at its building count, its beam schedule
    and its per-pour beam and drop feet), the material block (mixes, steel,
    PT, sand, the waste factors, the cartons' N), the labor, supervision,
    equipment and contract blocks with their costs, the PRICE row's sale and
    COST +, the tax flag;
  * the apply: the project and the estimate, the price sheet at the
    workbook's numbers (a mix the catalog has no price for gets its row),
    one section per tab at the tab's waste and markup with its labor subbed,
    the rows and the beam usages, the labor rates as section rates, the
    supervision and machine days typed onto their lines, a service the tab
    charges nothing for switched off, a machine the tab does not carry
    switched off, the cartons off, the misc items matched to the library or
    added; every section priced; a second import refused without --replace
    and rebuilt with it;
  * the tie-out report names the sections and the gaps.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import select, text

from app.models.estimate import Estimate
from app.models.estimate_labor import EstimateLaborLine
from app.models.estimate_section import EstimateSection
from app.models.grade_beam import GradeBeam
from app.models.misc_item import MiscItem
from app.models.mono_slab import MonoSlab
from app.models.project import Project
from app.services.workbook_import import (
    JobSpec,
    apply,
    app_mix_code,
    describe,
    read_workbook,
    tie_out,
)

D = Decimal


def _fill(ws, cells: dict[str, object]) -> None:
    for ref, value in cells.items():
        ws[ref] = value


def _slab_tab(ws) -> None:
    _fill(ws, {
        "B2": "JOB NAME", "B3": "Test Import", "G8": "Concrete Supplier", "J8": "Martin Marietta",
        "B9": "TYPE", "C9": "BLD #", "D9": "SQUARE FOOTAGE", "F9": "THICK INCH", "G9": "QTY ", "H9": "CABLE",
        "I9": "MIX DESIGN", "J9": "INCH OF SAND", "K9": "PERM. EDGE", "L9": "REINFORCING", "N9": "WIRE MESH Y/N",
        "O9": "EXP GB", "Q9": "Drops", "S9": "GRADE BEAMS", "AC9": "LABOR ADD", "AC10": "/SF",
        "G10": "SLABS", "L10": "SIZE", "M10": "SPACE", "O10": "TYPE", "P10": "FF", "Q10": "TYPE", "R10": "FF",
        "S10": "TYPE", "T10": "LN FT", "U10": "TYPE", "V10": "LN FT", "W10": "TYPE", "X10": "LN FT",
        # two building types: five T1s with two beams and a drop, seven T2s with one beam
        "B11": "01", "C11": "T1 / T1A", "D11": 7146, "F11": 4, "G11": 5, "H11": "y", "I11": 1, "J11": 2, "K11": 433,
        "N11": "N", "Q11": 6, "R11": 367, "S11": 1, "T11": 433, "U11": 2, "V11": 1251,
        "B12": "02", "C12": "T2", "D12": 4306, "F12": 4, "G12": 7, "H12": "y", "I12": 1, "J12": 2, "K12": 288,
        "S12": 1, "T12": 288, "AC12": 0.5,
        "O13": 9, "Q13": 10, "S13": 1,                       # a template placeholder row: no SF, no type
        "C49": "Total SF", "D49": 65872,
        "B50": "GB #", "C50": "GB DEMINSIONS", "E50": "TOP BARS", "G50": "BOT BARS", "I50": "MID BARS", "K50": "STIRRUPS",
        "C51": "WIDTH IN.", "D51": "HEIGHT IN", "E51": "# BARS", "F51": "SIZE", "G51": "# BARS", "H51": "SIZE",
        "I51": "# BARS", "J51": "SIZE", "K51": "SIZE", "L51": "SPACING",
        "B52": 1, "C52": 12, "D52": 26, "E52": 2, "F52": 5, "G52": 2, "H52": 5, "K52": 3, "L52": 24,
        "B53": 2, "C53": 12, "D53": 26, "E53": 2, "F53": 5, "K53": 3, "L53": 16,
        "B54": 3, "B55": 4, "B56": 5,
        "B57": 6, "C57": 3, "D57": 6, "E57": 1, "F57": 3, "K57": 3, "L57": 12,
        "B63": "UNIT SALE / FT", "E63": "SALE", "G63": "MARK UP OF SALE:", "J63": "UNIT COST:", "M63": "COST +",
        "B64": "PRICE:", "C64": 15.31, "E64": 1008500.5, "G64": 0.187, "J64": 12.45, "M64": 0.2,
        "Z63": "TAX  @…………", "AB63": "EXEMPT", "AC63": "n", "AD63": "Rate", "AE63": 0.0825,
        "B70": "Mix #", "C70": " MATERIAL COSTS :",
        "B71": 1, "C71": "3000 PSI W/ ASH PIERS, SOG", "G71": 155, "H71": "/CU. YD.", "I71": "CONCRETE WASTAGE", "K71": 0.08,
        "B72": 2, "C72": "3500 Psi", "G72": 160, "H72": "/CU. YD.",
        "B73": 3, "C73": "4000 psi", "G73": 165, "H73": "/CU. YD.",
        "B74": "TOTAL STEEL COST @", "G74": 0.7, "H74": "/LB", "I74": " WASTAGE", "K74": 0.1,
        "B75": "WIRE MESH", "G75": 0.26, "H75": "/SQ FT", "I75": " WASTAGE", "K75": 0.1,
        "B76": "POSTE TENSION", "G76": 0.85, "H76": "/SQ FT", "I76": "Quote:",
        "B77": " SAND", "G77": 25, "H77": "/YRD", "I77": "WASTAGE", "K77": 0.05,
        "B78": "CARTON  GRADE BEAMS", "E78": "Y/N         ?", "F78": "N", "G78": 1, "H78": "LN FT",
        "B82": " LABOR: SLAB", "D82": "Sub", "L82": "LABOR", "N82": "SUB LABOR",
        "B83": "FORMING", "D83": "y", "E83": 0.6, "G83": "/SQ FT.", "I83": 65872, "K83": "SQ FT", "L83": 0, "O83": 39523.2,
        "B84": "GRADING/CABLES", "D84": "y", "E84": 0.6, "G84": "/SQ FT", "L84": 0, "O84": 39523.2,
        "B85": "PLACE AND FINISH", "D85": "y", "E85": 0.65, "G85": "/SQ FT.", "L85": 0, "O85": 42816.8,
        "B86": "WRECK AND CLEAN UP", "D86": "y", "E86": 0.1, "G86": "/SQ FT.", "L86": 0, "O86": 6587.2,
        "B87": "DROPS", "D87": "y", "E87": 20, "G87": "/FC FT", "I87": 1835, "K87": "SQ FF", "L87": 0, "O87": 36700,
        "B88": "CABLE STRESSING", "D88": "y", "E88": 4, "G88": "EACH", "K88": "QTY", "L88": 0, "O88": 0,
        "B89": "TIE STEEL", "D89": "y", "E89": 400, "G89": "/TON", "I89": 30, "K89": "TONS =   ", "L89": 0, "O89": 12000,
        "B90": "COST PER SQ FT", "E90": 2.7, "I90": "TOTAL COST", "L90": 0, "O90": 177150.4,
        "B91": "SUPERVISION ",
        "B92": "Foreman", "D92": 16, "E92": 80, "F92": "DAYS@", "G92": 295, "H92": "/DAY", "O92": 23600,
        "B93": "Superintendent", "E93": 40, "F93": "DAYS@", "G93": 390, "H93": "/DAY", "O93": 15600,
        "B95": "Project management", "E95": 40, "F95": "DAYS@", "G95": 200, "O95": 8000,
        "B96": "EXPENSE ALLOWANCE", "E96": 80, "F96": "DAYS@", "G96": 100, "H96": "/DAY", "O96": 8000,
        "B97": "EQUIPMENT",
        "B98": "BOB CATS", "E98": 120, "F98": "Days ", "G98": 350, "H98": "/DAY", "O98": 19939.5,
        "B100": "TRENCHER", "E100": 120, "F100": "Days ", "G100": 300, "H100": "/DAY", "J100": 0.5, "O100": 17091,
        "B103": "MISCELLANIOUS", "E103": 120, "F103": "Days", "G103": 55, "H103": "/DAY", "O103": 1980,
        "B104": "CONTRACT  SERVICES",
        "B106": "SAW CUTTING", "F106": "/LN FT", "G106": 13174, "I106": "LN FT", "O106": 0,
        "B107": "CONCRETE PUMPING", "E107": 16, "F107": "/YRD", "G107": 2382, "I107": "CU YDS", "O107": 38112,
        "B108": "HAUL OFF", "E108": 4, "F108": "/YD", "I108": "YDS", "O108": 0,
        "B110": "OUT OF TOWN EXP", "F110": "DAYS", "G110": 250, "I110": "/DAY", "O110": 0,
        "B111": "MISCELLANIOUS", "G111": 1000, "I111": "FOR ", "O111": 0,
        "B115": "PRICE:", "C115": 15.31, "E115": 1008500.5, "M115": 0.2,
    })


def _paving_tab(ws) -> None:
    _fill(ws, {
        "B9": "PAVING TYPE", "D9": "SQUARE FOOTAGE", "F9": "THICK INCHES", "G9": "HAND POUR Y/N", "H9": "Traffic Control  Y/N",
        "I9": "LN FT DEMO", "J9": "LF CURBS", "K9": "MIX DESIGN", "L9": "LN FT THICK EDGE", "M9": "SAND INCHE",
        "N9": "STEEL REINFORCING", "P9": "MESH GAUGE", "N10": "SIZE", "O10": "SPACING",
        "B11": "Common medium duty", "D11": 47867, "F11": 7, "G11": "N", "H11": "N", "J11": 1882.7, "K11": 3, "M11": 2, "N11": 3, "O11": 18,
        "B12": "Retail light duty", "D12": 73096, "F12": 5, "G12": "N", "H12": "Y", "J12": 2875, "K12": 3, "M12": 2, "N12": 3, "O12": 18,
        "D28": 120963, "F28": "TOTAL SQ FT",
        "B31": "UNIT SALE / FT", "E31": "SALE", "G31": "MARK UP OF SALE:", "M31": "COST +",
        "B32": "PRICE:", "C32": 7.38, "E32": 893000, "G32": 0.187, "M32": 0.2,
        "R31": "TAX  @…………", "T31": "EXEMPT", "U31": "n", "V31": "Rate", "W31": 0.0825,
        "B39": "Mix #", "C39": " MATERIAL COSTS :",
        "B40": 2, "C40": "3500 Psi", "G40": 160, "H40": "/CU. YD.", "I40": "CONCRETE WASTAGE", "K40": 0.08,
        "B41": 3, "C41": "4000 psi", "G41": 165, "H41": "/CU. YD.",
        "C43": "TOTAL STEEL COST @", "G43": 0.65, "H43": "/LB", "I43": " WASTAGE", "K43": 0.1,
        "C44": "WIRE MESH", "G44": 0.4, "H44": "/SQ FT",
        "C46": "TOTAL COST SAND", "G46": 25, "H46": "/ YD", "I46": " WASTAGE", "K46": 0.05,
        "B49": " LABOR: SLAB", "D49": "Sub", "L49": "LABOR", "N49": "SUB LABOR",
        "B50": "FORMING", "D50": "y", "E50": 0.25, "G50": "/SQ FT.", "I50": 120963, "K50": "SQ FT", "L50": 0, "O50": 30240.75,
        "B51": "PLACE AND FINISH", "D51": "Y", "E51": 0.55, "G51": "/SQ FT.", "L51": 0, "O51": 66529.65,
        "B52": "WRECK AND CLEAN UP", "D52": "Y", "E52": 0.2, "G52": "/SQ FT.", "L52": 0, "O52": 24192.6,
        "B54": "CURB  ", "D54": "Y", "E54": 2, "G54": "LN FT", "I54": 4757.7, "K54": "LN FT", "L54": 0, "O54": 9515.4,
        "B55": "COST PER SQ FT", "E55": 1.08, "I55": "TOTAL COST", "L55": 0, "O55": 130478.4,
        "B56": "SUPERVISION ",
        "B57": "Foreman", "E57": 24, "F57": "DAYS@", "G57": 295, "H57": "/DAY", "O57": 7080,
        "B58": "Superintendent", "E58": 12, "F58": "DAYS@", "G58": 390, "H58": "/DAY", "O58": 4680,
        "B60": "Project manament", "E60": 12, "F60": "DAYS@", "G60": 200, "O60": 2400,
        "B61": "EXPENSE ALLOWANCE", "E61": 24, "F61": "DAYS@", "G61": 100, "H61": "/DAY", "O61": 2400,
        "B62": "EQUIPMENT",
        "B64": "Skid steer", "E64": 90, "F64": "Days ", "G64": 350, "H64": "/DAY", "J64": "FUEL&MAINT.", "O64": 14954.6,
        "B69": "CONTRACT  SERVICES",
        "B72": "SAW CUTTING ", "E72": 0.5, "F72": "/LN FT", "G72": 14784, "I72": "LN FT", "O72": 7392,
        "B73": "CONCRETE PUMP", "E73": 16, "F73": "/ YD", "G73": 2598, "I73": "YDS", "M73": "Pump Paving?", "N73": "N", "O73": 0,
        "B75": "CURB DEMO", "E75": 10, "F75": "/LN FT", "G75": 0, "I75": "LN FT", "O75": 0,
        "B77": "OUT OF TOWN EXP", "E77": 0, "F77": "DAYS", "G77": 250, "I77": "/DAY", "O77": 0,
        "B82": "PRICE:", "C82": 7.38, "E82": 893000, "M82": 0.2,
    })


def _footings_tab(ws) -> None:
    _fill(ws, {
        "E9": "MIX DESIGN", "F9": "Footing Deminsions", "I9": "Top Mat", "L9": "Bottom Mat", "T9": "PHILASTERS",
        "B10": "Type", "D10": "QTY", "F10": "LENGTH FT.", "G10": "Feet Wide", "H10": "INCHES THICK", "I10": "SPACING",
        "J10": "SIZE", "K10": "#  OF MATS", "L10": "SPACING", "M10": "SIZE", "N10": "# OF MATS", "T10": "QTY",
        "B11": "F8686 Retail", "D11": 13, "E11": 1, "F11": 8.5, "G11": 8.5, "H11": 20, "I11": 12, "J11": 6, "K11": 1, "L11": 12, "M11": 6, "N11": 1,
        "B12": "F5", "E12": 1,
        "E37": 13, "F37": "TOTAL LINEAR FEET BEAMS",
        "B40": "UNIT SALE / FT", "E40": "SALE", "G40": "MARK UP OF SALE:", "M40": "COST +",
        "B41": "PRICE:", "C41": 2700, "E41": 35100, "G41": 0.187, "M41": 0.2,
        "B50": "Mix #", "B51": 1, "C51": "3000 PSI W/ ASH PIERS, SOG", "G51": 155, "H51": "/CU. YD.", "I51": "CONCRETE WASTAGE", "K51": 0.1,
        "C54": "TOTAL STEEL COST @", "G54": 0.7, "H54": "/LB", "I54": " WASTAGE", "K54": 0.1,
        "B63": " LABOR: SLAB", "D63": "Sub", "L63": "LABOR", "N63": "SUB LABOR",
        "B65": "FORMING", "D65": "Y", "E65": 4, "G65": "/SQ FT.", "I65": 939, "K65": "FC.FT.=", "L65": 0, "O65": 3756,
        "B69": "TIE STEEL", "D69": "Y", "E69": 400, "G69": "/TON", "I69": 3.1, "K69": "TONS =   ", "L69": 0, "O69": 1240,
        "B70": "EXCAVATE / Backfill", "D70": "Y", "E70": 20, "G70": "CU YD", "I70": 64, "K70": "CU/YD", "L70": 0, "O70": 1280,
        "B71": "COST PER SQ FT", "E71": 6.7, "I71": "TOTAL COST", "L71": 0, "O71": 6276,
        "B72": "SUPERVISION MONO SLAB",
        "B73": "Foreman", "E73": 5, "F73": "DAYS@", "G73": 295, "H73": "/DAY", "O73": 1475,
        "B74": "Superintendent", "E74": 2.5, "F74": "DAYS@", "G74": 390, "H74": "/DAY", "O74": 975,
        "B77": "EXPENSE ALLOWANCE", "E77": 5, "F77": "DAYS@", "G77": 100, "H77": "/DAY", "O77": 500,
        "B78": "EQUIPMENT",
        "B84": "MISCELLANIOUS", "E84": 7, "F84": "Days ", "G84": 35, "H84": "$/DAY", "O84": 105,
        "B85": "CONTRACT  SERVICES",
        "B87": "PUMP ALLOWANCE", "E87": 16, "F87": "/YARD", "G87": 64, "I87": "YDS", "O87": 1024,
        "B94": "PRICE:", "C94": 2700, "E94": 35100, "M94": 0.2,
    })


def _misc_tab(ws) -> None:
    _fill(ws, {
        "A8": "SITE CONCRETE MIX MIX DESIGN:", "D8": 6, "H8": "3000 PSI Sidewalk and Hardscape", "K8": 155,
        "L8": "STRUCTURAL CONCRETE MIX DESIGN:", "O8": 6, "P8": "3000 PSI Sidewalk and Hardscape", "S8": 155,
        "B9": "TYPE", "E9": "SUB LABOR ?", "F9": "QTY", "G9": "UNIT SALE", "H9": "LABOR  COST", "I9": "TOTAL SALES",
        "J9": "DIA", "K9": "DEPTH", "M9": "YDS CONC.", "N9": "STEEL WT", "R9": "LABOR",
        "A10": "LIGHT POLE BASES", "E10": "Y", "F10": 25, "G10": 1250, "H10": 400, "I10": 31250, "J10": 16, "K10": 5.83, "M10": 8, "N10": 1749, "R10": 10000,
        "A11": "PIPE BOLLARDS", "E11": "Y", "G11": 250, "H11": 50,
        "A24": "AC CONDENSER PADS (1 per unit)", "E24": "Y", "F24": 10, "G24": 200, "H24": 80, "I24": 2000, "J24": 9, "K24": 4, "M24": 1.1, "N24": 30, "R24": 800,
    })


def build_workbook(path: Path) -> Path:
    wb = Workbook()
    info = wb.active
    info.title = "Information"
    _fill(info, {
        "B2": "JOB NAME:", "C2": "Test Import", "H2": " BID: ", "I2": "26-999", "J2": date(2026, 9, 9),
        "B3": "JOB ADDRESS:", "C3": "1 Main St", "H3": "REVISION: ", "I3": "GMP set",
        "B4": "CITY, STATE:", "C4": "Fort Worth, TX", "B8": "PREPARED BY:", "C8": "Chad McKinney",
        "B13": "CLIENT # 1 NAME:", "C13": "Test GC", "B16": "CONTACT:", "C16": "Somebody",
    })
    summary = wb.create_sheet("Summary")
    _fill(summary, {"O65": "Tax Exempt", "P65": "n", "Q65": "Tax Rate", "R65": 0.0825})
    pricing = wb.create_sheet("Pricing")
    _fill(pricing, {"M2": "LUMBER AND ACCSESS", "N2": "Unit Cost", "O2": "Unit", "M3": "2 X 4  X 16'", "N3": 0.99, "O3": "/LNFT",
                    "M4": "3/4 \" FORMING PLY", "N4": 75, "O4": "/SHEET"})
    _slab_tab(wb.create_sheet("Mono Slab on Grade Garden Style"))
    _paving_tab(wb.create_sheet("Private Paving"))
    wb.create_sheet("ROW Paving")            # an empty tab: skipped, no priced rows
    _footings_tab(wb.create_sheet("Footings"))
    _misc_tab(wb.create_sheet("Miscellaneous"))
    wb.create_sheet("PIERS")                 # a kind with no reader yet
    wb.create_sheet("Check List")
    wb.save(path)
    return path


@pytest.fixture
def workbook(tmp_path) -> Path:
    return build_workbook(tmp_path / "Test Import - Estimate - 26-999 - 2026-09-09_01.xlsm")


@pytest.fixture
def spec(workbook) -> JobSpec:
    return read_workbook(workbook)


# --------------------------------------------------------------- reading ----


def test_the_book_is_read_by_its_headers(spec):
    assert spec.project["name"] == "Test Import" and spec.project["job_number"] == "26-999"
    assert spec.project["gc"] == "Test GC" and spec.project["location"] == "1 Main St, Fort Worth, TX"
    assert spec.project["bid_date"] == date(2026, 9, 9) and spec.project["tax_exempt"] is False
    assert spec.estimate["name"] == "GMP set"
    assert spec.lumber == {"2 X 4 X 16'": D("0.99"), '3/4 " FORMING PLY': D("75")}
    assert spec.day_rates == {"labor_foreman_day_rate": D("295"), "labor_super_day_rate": D("390"),
                              "labor_pm_day_rate": D("200"), "labor_expense_day_rate": D("100")}
    assert [t.kind for t in spec.tabs] == ["mono_slab", "paving", "spot_footings", "miscellaneous"]
    assert dict(spec.skipped) == {"ROW Paving": "no priced rows", "PIERS": "a piers tab; no reader for it yet"}

    slab = spec.tabs[0]
    assert slab.name == "Mono Slab on Grade Garden Style" and slab.unit == "SF"
    assert [r["description"] for r in slab.rows] == ["01 T1 / T1A", "02 T2"]    # the placeholder row is not a pour
    t1 = slab.rows[0]
    assert (t1["square_footage"], t1["thickness_in"], t1["qty"], t1["post_tension"]) == (D("7146"), D("4"), 5, True)
    assert (t1["mix"], t1["sand_thickness_in"], t1["perimeter_edge_lf"], t1["wire_mesh"]) == (1, D("2"), D("433"), False)
    assert slab.rows[1]["qty"] == 7
    assert slab.rows[0]["paving_add_per_sf"] is None and slab.rows[1]["paving_add_per_sf"] == D("0.5")
    assert slab.usages == {0: [("grade_beam", 1, D("433")), ("grade_beam", 2, D("1251")), ("drop", 6, D("367"))],
                           1: [("grade_beam", 1, D("288"))]}
    assert [(b["n"], b["width_in"], b["height_in"], b["top_bars_count"], b["top_bars_size"], b["stirrup_spacing_in"])
            for b in slab.beam_types] == [(1, D("12"), D("26"), 2, 5, D("24")), (2, D("12"), D("26"), 2, 5, D("16")),
                                          (6, D("3"), D("6"), 1, 3, D("12"))]
    assert slab.quantity == D("65872") and slab.sale == D("1008500.5") and slab.margin == D("0.2")
    assert slab.total_cost == D("840417.08") and slab.tax_exempt is False and slab.supplier == "Martin Marietta"
    assert slab.mixes == {1: ("3000 PSI W/ ASH PIERS, SOG", D("155")), 2: ("3500 Psi", D("160")), 3: ("4000 psi", D("165"))}
    assert slab.prices == {"steel_lb": D("0.7"), "mesh_sf": D("0.26"), "pt_sf": D("0.85"), "sand_cy": D("25")}
    assert slab.waste == {"concrete": D("0.08"), "rebar": D("0.1"), "sand": D("0.05")}
    assert slab.switches == {"carton_forms": False}

    labor = {ln["label"]: ln for ln in slab.labor}
    assert labor["FORMING"]["rate"] == D("0.6") and labor["FORMING"]["sub"] is True and labor["FORMING"]["cost"] == D("39523.2")
    assert labor["DROPS"]["qty"] == D("1835") and labor["CABLE STRESSING"]["cost"] == 0
    assert [(s["label"], s["days"], s["rate"]) for s in slab.supervision] == [
        ("Foreman", D("80"), D("295")), ("Superintendent", D("40"), D("390")),
        ("Project management", D("40"), D("200")), ("EXPENSE ALLOWANCE", D("80"), D("100"))]
    assert [(e["label"], e["days"], e["rate"], e["cost"]) for e in slab.equipment] == [
        ("BOB CATS", D("120"), D("350"), D("19939.5")), ("TRENCHER", D("120"), D("300"), D("17091")),
        ("MISCELLANIOUS", D("120"), D("55"), D("1980"))]
    contract = {c["label"]: c for c in slab.contract}
    assert contract["CONCRETE PUMPING"]["rate"] == D("16") and contract["CONCRETE PUMPING"]["qty"] == D("2382")
    assert contract["SAW CUTTING"]["rate"] is None and contract["SAW CUTTING"]["cost"] == 0
    assert contract["OUT OF TOWN EXP"]["rate"] == D("250") and contract["OUT OF TOWN EXP"]["qty"] is None

    paving = spec.tabs[1]
    assert [r["description"] for r in paving.rows] == ["Common medium duty", "Retail light duty"]
    assert paving.rows[1]["traffic_control"] is True and paving.rows[0]["curb_lf"] == D("1882.7")
    assert paving.rows[0]["slab_bar_size"] == 3 and paving.rows[0]["slab_bar_spacing_in"] == D("18")
    pump = next(c for c in paving.contract if c["label"] == "CONCRETE PUMP")
    assert pump["enabled"] is False                                       # "Pump Paving?  N"
    assert next(c for c in paving.contract if c["label"].strip() == "SAW CUTTING")["cost"] == D("7392")

    footings = spec.tabs[2]
    assert len(footings.rows) == 1 and footings.quantity == D("13")
    f = footings.rows[0]
    assert (f["footing_count"], f["footing_each_ft"], f["ftg_width_in"], f["ftg_thick_in"]) == (13, D("8.5"), D("102.000"), D("20"))
    assert (f["ftg_top_spacing_in"], f["ftg_top_size"], f["ftg_bot_spacing_in"], f["ftg_bot_size"]) == (D("12"), 6, D("12"), 6)
    assert next(c for c in footings.contract if c["label"] == "PUMP ALLOWANCE")["unit"] == "/YARD"

    misc = spec.tabs[3]
    assert [(r["description"], r["qty"], r["unit_sale"], r["labor_per_unit"]) for r in misc.rows] == [
        ("LIGHT POLE BASES", D("25"), D("1250"), D("400")), ("AC CONDENSER PADS (1 per unit)", D("10"), D("200"), D("80"))]
    assert misc.rows[0]["mix"] == 6 and misc.mixes == {6: ("3000 PSI Sidewalk and Hardscape", D("155"))}
    assert misc.sale == D("33250")
    assert "Mono Slab on Grade Garden Style -> mono_slab" in describe(spec)


def test_a_mix_name_is_read_for_its_strength_and_its_ash():
    assert app_mix_code("3000 PSI W/ ASH PIERS, SOG") == "3000-ASH"
    assert app_mix_code("3500 Psi") == "3500-SC" and app_mix_code("4000 psi") == "4000-SC"
    assert app_mix_code("3000 PSI Sidewalk and Hardscape") == "3000-SC" and app_mix_code("no strength") is None


# ---------------------------------------------------------------- apply ----


def _lines(db, section_id):
    return {r.code: r for r in db.scalars(select(EstimateLaborLine).where(EstimateLaborLine.section_id == section_id)).all()}


def test_the_book_becomes_a_priced_job(db, spec):
    report = apply(db, spec)
    project = db.get(Project, report.project_id)
    estimate = db.get(Estimate, report.estimate_id)
    assert project.name == "Test Import" and project.job_number == "26-999" and project.gc == "Test GC"
    assert project.tax_exempt is False and project.bid_date == date(2026, 9, 9)
    assert estimate.name == "GMP set" and estimate.margin_pct == D("0.2") and estimate.contingency_pct == 0
    sections = list(db.scalars(select(EstimateSection).where(EstimateSection.estimate_id == estimate.id)
                               .order_by(EstimateSection.sort_order)).all())
    assert [s.kind for s in sections] == ["mono_slab", "paving", "spot_footings", "miscellaneous"]
    assert [s.name for s in sections] == ["Mono Slab on Grade Garden Style", "Private Paving", "Footings", "Miscellaneous"]
    slab, paving, footings, misc = sections
    assert slab.labor_subcontracted is True and slab.margin_pct == D("0.2") and slab.contingency_pct == 0
    assert (slab.waste_concrete, slab.waste_rebar, slab.waste_sand) == (D("0.08"), D("0.1"), D("0.05"))
    assert footings.waste_concrete == D("0.1") and footings.unit == "EA" and paving.unit == "SF"

    # The pours, at their building counts, with their beams.
    pours = list(db.scalars(select(MonoSlab).where(MonoSlab.section_id == slab.id).order_by(MonoSlab.sort_order)).all())
    assert [(p.description, p.qty, p.post_tension) for p in pours] == [("01 T1 / T1A", 5, True), ("02 T2", 7, True)]
    usages = db.scalars(select(GradeBeam).where(GradeBeam.mono_slab_id == pours[0].id).order_by(GradeBeam.sort_order)).all()
    assert [(u.beam_type.label, u.beam_type.kind, u.length_lf) for u in usages] == [
        ("GB 1", "grade_beam", D("433.000")), ("GB 2", "grade_beam", D("1251.000")), ("GB 6 (drop)", "drop", D("367.000"))]
    assert slab.calc_quantity == D("65872.000") and pours[0].calc_concrete_cy > 0 and pours[0].mix_design_id is not None

    # The job's price sheet: the mixes at the workbook's numbers, a row made where the catalog had none.
    sheet = {(r.kind, r.ref_id, r.ref_key): r for r in db.execute(text(
        "SELECT kind, ref_id, ref_key, value, is_edited, label FROM estimate_prices WHERE estimate_id = :e"), {"e": str(estimate.id)}).all()}
    mixes = {code: mid for code, mid in db.execute(text("SELECT code, id FROM mix_designs")).all()}
    assert sheet[("mix", mixes["3000-ASH"], None)].value == D("155.0000") and sheet[("mix", mixes["3000-ASH"], None)].is_edited
    assert sheet[("mix", mixes["4000-SC"], None)].value == D("165.0000")
    assert sheet[("setting", None, "labor_super_day_rate")].value == D("390.0000")
    lumber = next(r for r in sheet.values() if r.kind == "material" and r.label == "2 X 4  X 16'")
    assert lumber.value == D("0.9900")
    rebar = next(r for r in sheet.values() if r.kind == "material" and r.label == "REBAR PIERS / PT slabs")
    assert rebar.value == D("0.7000")

    # The tab's rates on the section; the days typed; what the tab does not carry switched off.
    rates = {k: D(str(v)) for k, v in db.execute(text("SELECT key, value FROM section_rates WHERE section_id = :s"), {"s": str(slab.id)}).all()}
    assert rates["labor_forming_sf"] == D("0.6") and rates["labor_drops_ff"] == D("20") and rates["labor_tie_steel_ton"] == D("400")
    assert rates["vapor_barrier_enabled"] == 0
    lines = _lines(db, slab.id)
    assert (lines["superintendent"].qty, lines["superintendent"].is_manual, lines["superintendent"].rate) == (D("40.0000"), True, D("390.0000"))
    assert (lines["foreman"].qty, lines["pm"].qty, lines["expense"].qty) == (D("80.0000"), D("40.0000"), D("80.0000"))
    assert lines["forming"].rate == D("0.6000") and lines["forming"].enabled and lines["forming"].qty == D("65872.0000")
    assert lines["drops"].qty == D("1835.0000") and lines["drops"].rate == D("20.0000")     # 367 FF x 5 buildings
    assert lines["labor_add"].qty == D("15071.0000") and lines["labor_add"].ext_cost == D("15071.00")   # 4306 SF x 7 x $0.50
    assert pours[1].paving_add_per_sf == D("0.5000") and pours[0].paving_add_per_sf is None
    equipment = {r.code: r for r in db.execute(text(
        "SELECT code, days_qty, rate, enabled, is_manual, ext_cost FROM estimate_equipment_lines WHERE section_id = :s"), {"s": str(slab.id)}).all()}
    assert (equipment["trencher"].days_qty, equipment["trencher"].rate, equipment["trencher"].is_manual) == (D("120.0000"), D("300.0000"), True)
    assert equipment["skid_steer"].days_qty == D("120.0000") and equipment["skid_steer"].rate == D("350.0000")   # the tab's BOB CATS
    assert equipment["mini_excavator"].enabled is False                                    # not on the tab
    assert equipment["haul_off"].ext_cost == 0 and equipment["concrete_pump"].enabled       # charged nothing / charged
    slab_entry = report.sections[0]
    assert any(o.startswith("mini_excavator (") for o in slab_entry["off"])   # the tab has no mini excavator

    # Paving: the pump the tab turned off, the traffic control, the saw as the app's soft cut.
    paving_lines = {r.code: r for r in db.execute(text(
        "SELECT code, enabled, days_qty FROM estimate_equipment_lines WHERE section_id = :s"), {"s": str(paving.id)}).all()}
    assert paving_lines["concrete_pump"].enabled is False
    prates = {k: D(str(v)) for k, v in db.execute(text("SELECT key, value FROM section_rates WHERE section_id = :s"), {"s": str(paving.id)}).all()}
    assert prates["labor_curb_lf"] == D("2") and prates["joint_soft_cut_lf"] == D("0.5")
    areas = list(db.scalars(select(MonoSlab).where(MonoSlab.section_id == paving.id).order_by(MonoSlab.sort_order)).all())
    assert areas[1].traffic_control is True and areas[0].curb_lf == D("1882.700")

    # Footings: a spot footing at its count; the yards unit read as CY.
    run = db.execute(text("SELECT footing_count, footing_each_ft, length_ft, ftg_width_in, ftg_thick_in, ftg_top_size FROM wall_runs WHERE section_id = :s"),
                     {"s": str(footings.id)}).one()
    assert (run.footing_count, run.footing_each_ft, run.length_ft, run.ftg_width_in, run.ftg_thick_in, run.ftg_top_size) == (13, D("8.500"), D("110.500"), D("102.000"), D("20.000"), 6)
    frates = {k: D(str(v)) for k, v in db.execute(text("SELECT key, value FROM section_rates WHERE section_id = :s"), {"s": str(footings.id)}).all()}
    assert frates["concrete_pump_cy"] == D("16") and frates["labor_excavate_cy"] == D("20")

    # Misc: the library item at the tab's count and dimensions, the pad added as a slab.
    items = {i.description: i for i in db.scalars(select(MiscItem).where(MiscItem.section_id == misc.id)).all()}
    poles = items["Light Pole Bases"]
    assert (poles.qty, poles.unit_sale, poles.labor_per_unit, poles.dim_a, poles.dim_b) == (D("25.000"), D("1250.00"), D("400.00"), D("16.000"), D("5.830"))
    pad = items["AC CONDENSER PADS (1 per unit)"]
    assert (pad.shape, pad.qty, pad.dim_a, pad.dim_b, pad.unit_sale) == ("slab", D("10.000"), D("9.000"), D("4.000"), D("200.00"))
    assert pad.calc_concrete_cy == (D("9") * D("4") / D("324") * D("10")).quantize(D("0.0001"))
    assert all(i.qty == 0 for name, i in items.items() if name not in ("Light Pole Bases", "AC CONDENSER PADS (1 per unit)"))

    # Every section priced, the job rolled up, the tie-out naming the sections.
    for s in sections:
        assert s.calc_total_sale and s.calc_total_sale > 0, s.name
    assert misc.calc_total_sale == D("33250.00")
    db.refresh(estimate)
    assert estimate.calc_total_sale == sum(s.calc_total_sale for s in sections)
    out = tie_out(report)
    assert "Mono Slab on Grade Garden Style" in out and "TOTAL" in out and "Private Paving: 2 rows" in out
    assert report.tab_sale == D("1008500.5") + D("893000") + D("35100") + D("33250")
    assert report.app_sale == estimate.calc_total_sale


def test_a_second_import_is_refused_without_replace_and_rebuilt_with_it(db, spec):
    first = apply(db, spec)
    with pytest.raises(RuntimeError, match="already exists"):
        apply(db, spec)
    second = apply(db, spec, replace=True)
    assert second.project_id == first.project_id and second.estimate_id != first.estimate_id
    assert db.get(Estimate, first.estimate_id) is None
    assert db.scalar(select(Estimate).where(Estimate.project_id == first.project_id, Estimate.name == "GMP set")).id == second.estimate_id
    # Another name on the same project is another estimate.
    third = apply(db, spec, estimate_name="Rev 2")
    assert third.estimate_id not in (first.estimate_id, second.estimate_id)
    assert db.scalar(select(text("count(*)")).select_from(Estimate).where(Estimate.project_id == first.project_id)) == 2
